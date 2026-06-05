"""
Post-analysis Excel helpers: Run Info sheet, foci CSV integration, shared channel-row collection.
"""
from __future__ import annotations

import csv
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Optional

import numpy as np
import pandas as pd

from bactomate.gui.analysis_paths import base_dir_for_mask, stem_from_mask_path
from bactomate.gui.artifact_paths import (
    infer_pipeline_dir_from_mask,
    resolve_filtered_foci_csv,
    resolve_segmentation_foci_csv,
)
from bactomate.gui.utils import safe_imread

RUN_INFO_SHEET = "Run Info"
MASTER_PER_CELL_SUMMARY_BASENAME = "@summary_per-cell-data"
MASTER_PER_CELL_SUMMARY_FILENAME = f"{MASTER_PER_CELL_SUMMARY_BASENAME}.xlsx"
MASTER_PER_CELL_SUMMARY_SHEET = MASTER_PER_CELL_SUMMARY_BASENAME
MASTER_PER_CELL_SOURCE_COLUMN = "Analysis workbook"
MASTER_PER_CELL_SAMPLE_COLUMN = "Sample"
PER_CELL_DATA_SHEET = "Per-Cell Data"

# Summary columns managed by post-analysis foci / coloc updates (see apply_foci_summary_columns).
SUMMARY_FOCI_COL_PREFIXES = (
    "Total Foci Count (",
    "Total Foci Count within Cells (",  # legacy
    "Foci outside cells (",
)


@dataclass
class FociChannelArtifact:
    """One foci channel resolved for post-analysis import."""

    channel_label: str
    foci_stack: Path
    csv_path: Path
    csv_folder: str  # e.g. "foci_filtered" or "foci_detection"


@dataclass
class FociIntegrationOutcome:
    """Result of importing segmentation foci CSVs into an analysis workbook."""

    performed: bool = False
    channels: list[FociChannelArtifact] = field(default_factory=list)
    foci_rows_written: int = 0
    in_cell_foci: int = 0
    per_cell_columns: list[str] = field(default_factory=list)
    messages: list[str] = field(default_factory=list)


def _safe_col_key(raw: str) -> str:
    txt = str(raw or "").strip()
    if not txt:
        return "foci"
    out = []
    for ch in txt:
        out.append(ch if ch.isalnum() or ch == "_" else "_")
    key = "".join(out).strip("_")
    while "__" in key:
        key = key.replace("__", "_")
    return key or "foci"


def _read_foci_csv_rows(csv_path: Path) -> list[dict[str, Any]]:
    if csv_path is None or not csv_path.exists():
        return []
    rows: list[dict[str, Any]] = []
    with open(csv_path, "r", newline="") as f:
        reader = csv.DictReader(f)
        for r in reader:
            try:
                y = int(float(r.get("y", 0) or 0))
                x = int(float(r.get("x", 0) or 0))
            except Exception:
                continue
            try:
                intensity = float(r.get("intensity", 0) or 0.0)
            except Exception:
                intensity = 0.0
            prob_raw = r.get("probability", None)
            prob = None
            if prob_raw not in (None, "", "None"):
                try:
                    prob = float(prob_raw)
                except Exception:
                    prob = None
            frame = 0
            if "frame" in (reader.fieldnames or []):
                try:
                    frame = int(float(r.get("frame", 0) or 0))
                except Exception:
                    frame = 0
            try:
                size_px = int(float(r.get("area_px", r.get("size_px", 1)) or 1))
            except (TypeError, ValueError):
                size_px = 1
            mean_intensity = r.get("mean_intensity", None)
            rel_intensity = r.get("rel_intensity", None)
            cell_id_raw = r.get("cell_id", None)
            cell_id = None
            if cell_id_raw not in (None, "", "None"):
                try:
                    cell_id = int(float(cell_id_raw))
                except (TypeError, ValueError):
                    cell_id = None
            rows.append(
                {
                    "frame": frame,
                    "y": y,
                    "x": x,
                    "intensity": intensity,
                    "probability": prob,
                    "size_px": size_px,
                    "mean_intensity": mean_intensity,
                    "rel_intensity": rel_intensity,
                    "cell_id": cell_id,
                }
            )
    return rows


def _csv_path_for_source(
    pipeline_dir: Path,
    *,
    mask_path: Path,
    foci_source_path: Path,
    prefer_filtered: bool,
) -> tuple[Optional[Path], Optional[str]]:
    """Return (csv path, folder name) for the CSV actually used."""
    if prefer_filtered:
        filt = resolve_filtered_foci_csv(
            pipeline_dir, mask_path=mask_path, foci_source_path=foci_source_path
        )
        if filt is not None:
            return filt, filt.parent.name
    det = resolve_segmentation_foci_csv(
        pipeline_dir, mask_path=mask_path, foci_source_path=foci_source_path
    )
    if det is not None:
        return det, det.parent.name
    return None, None


def _channel_label_for_source(tab, mask_path: Path, foci_source_path: Path, used: set[str]) -> str:
    mask_stem = mask_path.stem.replace("_masks", "").replace("_mask_filtered", "")
    tag = None
    try:
        tag = tab._extract_channel_tag_for_mask_stem(Path(foci_source_path).stem, mask_stem)
    except Exception:
        tag = None
    base = (tag or "").strip() or "foci"
    lbl = base
    n = 2
    while lbl in used:
        lbl = f"{base}_{n}"
        n += 1
    used.add(lbl)
    return lbl


def _load_mask_frames(mask_path: Path) -> dict[int, np.ndarray]:
    try:
        arr = np.asarray(safe_imread(str(mask_path)))
        if arr.ndim == 2:
            return {0: arr.astype(np.int32)}
        if arr.ndim >= 3:
            out: dict[int, np.ndarray] = {}
            for i in range(int(arr.shape[0])):
                out[i] = np.asarray(arr[i], dtype=np.int32)
            return out
    except Exception:
        return {}
    return {}


def discover_foci_csv_artifacts_for_mask(
    tab,
    mask_path: Path,
    *,
    prefer_filtered: bool = True,
) -> list[FociChannelArtifact]:
    """Resolve foci stacks + CSV paths without reading the mask or assigning peaks."""
    mask_path = Path(mask_path)
    pipeline_dir = infer_pipeline_dir_from_mask(mask_path)
    try:
        sources = tab._resolve_foci_paths_for_mask(mask_path)
    except Exception:
        sources = []
    if not sources:
        return []

    used_labels: set[str] = set()
    used_csv_keys: set[str] = set()
    artifacts: list[FociChannelArtifact] = []

    for src in sources:
        csv_path, csv_folder = _csv_path_for_source(
            pipeline_dir, mask_path=mask_path, foci_source_path=src, prefer_filtered=prefer_filtered
        )
        if csv_path is None or not Path(csv_path).is_file():
            continue
        csv_key = str(Path(csv_path).resolve())
        if csv_key in used_csv_keys:
            continue
        used_csv_keys.add(csv_key)
        lbl = _channel_label_for_source(tab, mask_path, src, used_labels)
        artifacts.append(
            FociChannelArtifact(
                channel_label=lbl,
                foci_stack=Path(src),
                csv_path=Path(csv_path),
                csv_folder=str(csv_folder or Path(csv_path).parent.name),
            )
        )
    return artifacts


def collect_channel_rows_for_mask(
    tab,
    mask_path: Path,
    *,
    prefer_filtered: bool = True,
) -> tuple[
    list[tuple[str, list[dict[str, Any]]]],
    dict[tuple[str, int], int],
    list[str],
    list[FociChannelArtifact],
]:
    """
    Load foci CSVs for each matching *_foci* stack and assign peaks to cells.

    Returns:
        channel_rows: (label, assigned row dicts with cell_id, in_cell, frame, …)
        raw_by_label_frame: raw CSV row counts per (label, csv_frame_index)
        log_lines
        artifacts: resolved stack + CSV path per channel
    """
    from bactomate.gui.foci_filter_pipeline import (
        filter_foci_csv_rows,
        filter_foci_rows_by_min_area_px,
    )
    from bactomate.modules.foci_utils import assign_peaks_to_cells

    log_lines: list[str] = []
    mask_path = Path(mask_path)
    mask_frames = _load_mask_frames(mask_path)
    if not mask_frames:
        return [], {}, [f"[Foci] Skipped {mask_path.name}: could not read mask.\n"], []

    pipeline_dir = infer_pipeline_dir_from_mask(mask_path)
    try:
        sources = tab._resolve_foci_paths_for_mask(mask_path)
    except Exception:
        sources = []
    if not sources:
        return [], {}, [f"[Foci] Skipped {mask_path.name}: no matching *_foci image stacks.\n"], []

    used_labels: set[str] = set()
    used_csv_keys: set[str] = set()
    channel_rows: list[tuple[str, list[dict[str, Any]]]] = []
    raw_by_label_frame: dict[tuple[str, int], int] = defaultdict(int)
    artifacts: list[FociChannelArtifact] = []

    for src in sources:
        csv_path, csv_folder = _csv_path_for_source(
            pipeline_dir, mask_path=mask_path, foci_source_path=src, prefer_filtered=prefer_filtered
        )
        if csv_path is None:
            log_lines.append(f"[Foci] {Path(src).name} -> no foci CSV found.\n")
            continue
        csv_key = str(Path(csv_path).resolve())
        if csv_key in used_csv_keys:
            log_lines.append(
                f"[Foci] Skip {Path(src).name}: same CSV already imported ({Path(csv_path).name}).\n"
            )
            continue
        used_csv_keys.add(csv_key)

        lbl = _channel_label_for_source(tab, mask_path, src, used_labels)
        csv_rows = _read_foci_csv_rows(csv_path)
        if not csv_rows:
            log_lines.append(f"[Foci] {Path(src).name} -> empty CSV {csv_path.name}.\n")
            continue
        ch_settings: dict[str, Any] = {}
        if hasattr(tab, "_foci_filter_settings_for_source"):
            try:
                ch_settings = dict(tab._foci_filter_settings_for_source(Path(src), mask_path))
            except Exception:
                ch_settings = {}
        n_csv_raw = len(csv_rows)
        for r in csv_rows:
            try:
                fr_csv = int(r.get("frame", 0) or 0)
            except Exception:
                fr_csv = 0
            raw_by_label_frame[(lbl, fr_csv)] += 1

        assigned_all: list[dict[str, Any]] = []
        n_after_filter = 0
        max_mask_idx = max(mask_frames.keys()) if mask_frames else 0
        frames = sorted({int(r.get("frame", 0) or 0) for r in csv_rows})
        for fr in frames:
            fr_use = max(0, min(int(fr), int(max_mask_idx)))
            mask_labels = mask_frames.get(fr_use)
            if mask_labels is None:
                continue
            peaks = [r for r in csv_rows if int(r.get("frame", 0) or 0) == int(fr)]
            label_plane = np.asarray(mask_labels, dtype=np.int32)
            accepted_labels = {
                int(lbl) for lbl in np.unique(label_plane) if int(lbl) > 0
            }
            peaks = filter_foci_csv_rows(
                peaks,
                ch_settings,
                mask_2d=label_plane,
                accepted_cell_labels=accepted_labels,
            )
            peaks = filter_foci_rows_by_min_area_px(peaks, ch_settings)
            n_after_filter += len(peaks)
            foci_rows, _cell_map = assign_peaks_to_cells(
                peaks, label_plane, dilation_px=1, frame_index=int(fr)
            )
            for r in foci_rows:
                q = dict(r)
                q["channel_name"] = lbl
                q["frame"] = int(fr)
                assigned_all.append(q)
        channel_rows.append((lbl, assigned_all))
        artifacts.append(
            FociChannelArtifact(
                channel_label=lbl,
                foci_stack=Path(src),
                csv_path=Path(csv_path),
                csv_folder=str(csv_folder or csv_path.parent.name),
            )
        )
        log_lines.append(
            f"[Foci] Channel '{lbl}' from {Path(src).name}: csv_rows={n_csv_raw}, "
            f"after_filters={n_after_filter}, peaks_assigned={len(assigned_all)}, "
            f"csv={csv_path} ({csv_folder})\n"
        )

    return channel_rows, dict(raw_by_label_frame), log_lines, artifacts


def build_foci_table_from_channel_rows(
    channel_rows: list[tuple[str, list[dict[str, Any]]]],
) -> list[dict[str, Any]]:
    """Flatten channel rows into analyze_cell_areas-style foci_table records."""
    out: list[dict[str, Any]] = []
    for _lbl, rows in channel_rows:
        next_id = 1
        for r in rows or []:
            if not bool(r.get("in_cell", False)):
                continue
            cid = r.get("cell_id")
            if cid in (None, "", "None"):
                continue
            try:
                if int(cid) <= 0:
                    continue
            except Exception:
                continue
            intensity = r.get("intensity")
            size_px = r.get("size_px", 1)
            out.append(
                {
                    "foci_id": next_id,
                    "cell_id": int(cid),
                    "frame": int(r.get("frame", 0) or 0),
                    "channel_name": r.get("channel_name") or _lbl,
                    "y": int(r.get("y", 0)),
                    "x": int(r.get("x", 0)),
                    "intensity": intensity,
                    "size_px": size_px if size_px is not None else 1,
                    "mean_intensity": r.get("mean_intensity", intensity),
                    "rel_intensity": r.get("rel_intensity"),
                    "in_cell": True,
                }
            )
            next_id += 1
    return out


def _aggregate_foci_metrics_by_cell(
    foci_table: list[dict[str, Any]],
    *,
    frame_shift: int,
) -> dict[tuple[int, int, str], dict[str, Any]]:
    """Group in-cell foci by (summary_frame, cell_id, channel_label)."""
    grouped: dict[tuple[int, int, str], dict[str, Any]] = {}
    for r in foci_table:
        try:
            fr = int(r.get("frame", 0) or 0) + int(frame_shift)
            cid = int(r["cell_id"])
            lbl = str(r.get("channel_name") or "foci")
        except Exception:
            continue
        key = (fr, cid, lbl)
        bucket = grouped.setdefault(
            key,
            {"foci_ids": [], "intensities": [], "mean_intensities": [], "size_px": []},
        )
        fid = r.get("foci_id")
        if fid is not None:
            bucket["foci_ids"].append(int(fid))
        intensity = r.get("intensity")
        if intensity is not None:
            try:
                val = float(intensity)
                if np.isfinite(val):
                    bucket["intensities"].append(val)
            except Exception:
                pass
        mean_int = r.get("mean_intensity", intensity)
        if mean_int is not None:
            try:
                val = float(mean_int)
                if np.isfinite(val):
                    bucket["mean_intensities"].append(val)
            except Exception:
                pass
        size_px = r.get("size_px", 1)
        try:
            bucket["size_px"].append(float(size_px) if size_px is not None else 1.0)
        except Exception:
            bucket["size_px"].append(1.0)
    return grouped


def _sorted_foci_lists_for_cell(agg: dict[str, Any]) -> tuple[list[int], list[float], list[float], list[float]]:
    """Return foci_ids, peak intensities, mean intensities, sizes in matching foci_id order."""
    ids = list(agg.get("foci_ids") or [])
    ints = list(agg.get("intensities") or [])
    means = list(agg.get("mean_intensities") or [])
    sizes = list(agg.get("size_px") or [])
    n = len(ids)
    while len(ints) < n:
        ints.append(np.nan)
    while len(means) < n:
        means.append(ints[len(means)] if len(means) < len(ints) else np.nan)
    while len(sizes) < n:
        sizes.append(1.0)
    combined = sorted(zip(ids, ints, means, sizes), key=lambda t: int(t[0]))
    if not combined:
        return [], [], [], []
    ids_o, ints_o, means_o, sizes_o = zip(*combined)
    return list(ids_o), list(ints_o), list(means_o), list(sizes_o)


def _update_per_cell_foci_detail_columns(
    per_cell: pd.DataFrame,
    *,
    channel_rows: list[tuple[str, list[dict[str, Any]]]],
    foci_table: list[dict[str, Any]],
    frame_shift: int,
) -> tuple[pd.DataFrame, list[str]]:
    """
    Add per-cell foci columns per channel (from CSV, not FI re-detection).

    For each cell with multiple foci:
    - ``foci_ids_{ch}`` plus comma-separated per-focus lists (same order as IDs)
    - ``{ch}_foci_intensities``, ``{ch}_foci_mean_intensities``, ``{ch}_foci_sizes_px``
    - ``{ch}_foci_mean_intensity`` / ``_total_`` / ``_max_`` are **cell-level** summaries
      (average, sum, and max across that cell's foci).
    """
    cell_col = "Cell_ID" if "Cell_ID" in per_cell.columns else ("Cell Label" if "Cell Label" in per_cell.columns else None)
    if cell_col is None:
        return per_cell, []

    has_frame = "Frame" in per_cell.columns
    metrics = _aggregate_foci_metrics_by_cell(foci_table, frame_shift=frame_shift)
    added_cols: list[str] = []

    for lbl, _rows in channel_rows:
        safe = _safe_col_key(lbl)
        col_ids = f"foci_ids_{safe}"
        col_ints = f"{safe}_foci_intensities"
        col_mean_list = f"{safe}_foci_mean_intensities"
        col_sizes = f"{safe}_foci_sizes_px"
        col_mean = f"{safe}_foci_mean_intensity"
        col_total = f"{safe}_foci_total_intensity"
        col_max = f"{safe}_foci_max_intensity"
        text_cols = (col_ids, col_ints, col_mean_list, col_sizes)
        for col in (*text_cols, col_mean, col_total, col_max):
            if col not in per_cell.columns:
                per_cell[col] = "" if col in text_cols else np.nan
            if col not in added_cols:
                added_cols.append(col)

        def _row_values(fr_val: Any, cid_val: Any) -> tuple[str, str, str, str, float, float, float]:
            try:
                fr_i = int(fr_val) if fr_val is not None and pd.notna(fr_val) else 0
            except Exception:
                fr_i = 0
            try:
                cid_i = int(cid_val) if cid_val is not None and pd.notna(cid_val) else -1
            except Exception:
                cid_i = -1
            agg = metrics.get((fr_i, cid_i, lbl), {})
            ids, ints, means, sizes = _sorted_foci_lists_for_cell(agg)
            mean_vals = [m for m in means if m is not None and np.isfinite(m)]
            int_vals = [v for v in ints if v is not None and np.isfinite(v)]
            mean_i = float(np.mean(mean_vals if mean_vals else int_vals)) if (mean_vals or int_vals) else np.nan
            total_i = float(np.sum(int_vals)) if int_vals else np.nan
            max_i = float(np.max(int_vals)) if int_vals else np.nan
            ints_s = ", ".join(f"{v:.4g}" for v in ints) if ints else ""
            means_s = ", ".join(f"{v:.4g}" for v in means) if means else ""
            sizes_s = ", ".join(f"{v:.4g}" for v in sizes) if sizes else ""
            return (
                ",".join(str(x) for x in ids),
                ints_s,
                means_s,
                sizes_s,
                mean_i,
                total_i,
                max_i,
            )

        if has_frame:
            values = [_row_values(fr, cid) for fr, cid in zip(per_cell["Frame"], per_cell[cell_col])]
        else:
            values = [_row_values(0, cid) for cid in per_cell[cell_col]]

        per_cell[col_ids] = [v[0] for v in values]
        per_cell[col_ints] = [v[1] for v in values]
        per_cell[col_mean_list] = [v[2] for v in values]
        per_cell[col_sizes] = [v[3] for v in values]
        per_cell[col_mean] = [v[4] for v in values]
        per_cell[col_total] = [v[5] for v in values]
        per_cell[col_max] = [v[6] for v in values]

    return per_cell, added_cols


def apply_foci_summary_columns(
    summary: pd.DataFrame,
    *,
    label: str,
    totals_by_frame: dict[int, int],
    within_by_frame: dict[int, int],
) -> pd.DataFrame:
    """
    Add per-channel foci counts to the Summary sheet.

    ``Total Foci Count ({channel})`` is always the in-cell count (matches Per-Cell / Foci sheet).
    When the CSV still contains peaks outside cells, also add ``Foci outside cells ({channel})``.
    """
    col_total = f"Total Foci Count ({label})"
    col_outside = f"Foci outside cells ({label})"

    def _map_counts(counts: dict[int, int]) -> pd.Series:
        return summary["Frame"].map(
            lambda f: int(counts.get(int(f), 0)) if f is not None and pd.notna(f) else 0
        )

    summary[col_total] = _map_counts(within_by_frame)
    excluded: dict[int, int] = {}
    frames = set(totals_by_frame.keys()) | set(within_by_frame.keys())
    for fr in frames:
        excluded[fr] = max(0, int(totals_by_frame.get(fr, 0)) - int(within_by_frame.get(fr, 0)))
    if any(v > 0 for v in excluded.values()):
        summary[col_outside] = _map_counts(excluded)
    return summary


def outcome_to_manifest_entries(outcome: FociIntegrationOutcome) -> dict[str, Any]:
    """Map integration outcome to Run Info manifest keys."""
    entries: dict[str, Any] = {}
    if not outcome.channels and not outcome.performed:
        if outcome.messages:
            entries["Foci analysis result"] = "No foci stacks or CSVs found"
        return entries
    entries["Foci analysis performed"] = "Yes" if outcome.performed else "No"
    if outcome.channels:
        entries["Foci channels"] = ", ".join(a.channel_label for a in outcome.channels)
        csv_lines = [
            f"{a.channel_label}: {a.csv_path.resolve()} ({a.csv_folder})" for a in outcome.channels
        ]
        entries["Foci CSV files used"] = "\n".join(csv_lines)
    if outcome.in_cell_foci:
        entries["In-cell foci written"] = str(int(outcome.in_cell_foci))
    if outcome.per_cell_columns:
        entries["Per-Cell foci columns"] = ", ".join(outcome.per_cell_columns)
    return entries


def refresh_workbooks_after_filtering(
    tab,
    mask_paths: list[Path],
    *,
    prefer_filtered: bool = True,
) -> list[str]:
    """
    After Filter File/Folder, push filtered foci CSV metrics into any existing analysis workbooks.

    Uses the filtered mask for cell assignment when available. Skips masks with no workbook.
    """
    from bactomate.gui.analysis_paths import resolve_analysis_workbook_for_mask
    from bactomate.gui.artifact_paths import filtered_mask_path_for_mask

    log_lines: list[str] = []
    for mask_path in mask_paths:
        mask_path = Path(mask_path)
        workbook = resolve_analysis_workbook_for_mask(mask_path)
        if workbook is None:
            continue
        mask_for_cells = mask_path
        try:
            filt = filtered_mask_path_for_mask(mask_path)
            if filt.is_file():
                mask_for_cells = filt
        except Exception:
            pass
        try:
            outcome = integrate_foci_csv_into_workbook(
                tab,
                mask_for_cells,
                workbook,
                prefer_filtered=prefer_filtered,
            )
            log_lines.extend(outcome.messages)
        except Exception as exc:
            log_lines.append(f"[Filtering] Workbook foci refresh failed for {mask_path.name}: {exc}\n")
    try:
        from bactomate.gui.foci_overlay_artifacts import refresh_filtered_foci_overlays_after_filtering

        log_lines.extend(refresh_filtered_foci_overlays_after_filtering(tab, mask_paths))
    except Exception as exc:
        log_lines.append(f"[Filtering] Filtered overlay refresh failed: {exc}\n")
    return log_lines


def integrate_foci_csv_into_workbook(
    tab,
    mask_path: Path,
    analysis_path: Path,
    *,
    prefer_filtered: bool = True,
) -> FociIntegrationOutcome:
    """
    Integrate Spotiflow/detection foci CSVs into an existing analysis workbook.

    Writes/replaces the Foci sheet and updates Per-Cell / Summary count and intensity columns.
    Intensities come from the CSV peak values (independent of Analyze FI).
    """
    from bactomate.gui.tabs._analyze_masks_colocalization import (
        _infer_frame_shift,
        _persist_workbook_count_updates,
    )

    outcome = FociIntegrationOutcome()
    analysis_path = Path(analysis_path)
    if not analysis_path.is_file():
        outcome.messages.append(f"[Foci] Skipped: workbook not found ({analysis_path}).\n")
        return outcome

    channel_rows, raw_counts, logs, artifacts = collect_channel_rows_for_mask(
        tab, mask_path, prefer_filtered=prefer_filtered
    )
    outcome.messages.extend(logs)
    outcome.channels = artifacts

    if not channel_rows:
        return outcome

    foci_table = build_foci_table_from_channel_rows(channel_rows)
    outcome.in_cell_foci = len(foci_table)

    try:
        _persist_workbook_count_updates(
            analysis_path=analysis_path,
            channel_rows=channel_rows,
            raw_counts=raw_counts,
            pair_tables=[],
            same_cell_only=False,
        )
        for lbl, _rows in channel_rows:
            outcome.per_cell_columns.append(f"{_safe_col_key(lbl)}_foci_count")
    except Exception as exc:
        outcome.messages.append(f"[Foci] Warning: count columns not updated: {exc}\n")

    per_cell_columns_added: list[str] = []
    try:
        per_cell = pd.read_excel(analysis_path, sheet_name="Per-Cell Data")
    except Exception as exc:
        per_cell = None
        outcome.messages.append(f"[Foci] Warning: could not read Per-Cell Data: {exc}\n")

    if per_cell is not None and not per_cell.empty and channel_rows:
        try:
            has_frame = "Frame" in per_cell.columns
            per_cell_frames: set[int] = set()
            if has_frame:
                try:
                    per_cell_frames = {
                        int(v) for v in per_cell["Frame"].tolist() if v is not None and pd.notna(v)
                    }
                except Exception:
                    per_cell_frames = set()
            frame_shift = _infer_frame_shift(
                per_cell_frames=per_cell_frames,
                channel_rows=channel_rows,
                pair_tables=[],
            )
            if foci_table:
                per_cell, per_cell_columns_added = _update_per_cell_foci_detail_columns(
                    per_cell,
                    channel_rows=channel_rows,
                    foci_table=foci_table,
                    frame_shift=frame_shift,
                )
                outcome.per_cell_columns.extend(per_cell_columns_added)
            with pd.ExcelWriter(analysis_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                per_cell.to_excel(writer, sheet_name="Per-Cell Data", index=False)
        except Exception as exc:
            outcome.messages.append(f"[Foci] Warning: Per-Cell foci columns not updated: {exc}\n")

    if foci_table:
        df_foci = pd.DataFrame(foci_table)
        col_order = [
            "foci_id",
            "cell_id",
            "frame",
            "channel_name",
            "y",
            "x",
            "intensity",
            "size_px",
            "mean_intensity",
            "rel_intensity",
            "in_cell",
        ]
        df_foci = df_foci[[c for c in col_order if c in df_foci.columns]]
        with pd.ExcelWriter(analysis_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
            df_foci.to_excel(writer, sheet_name="Foci", index=False)
        outcome.foci_rows_written = len(foci_table)
        outcome.performed = True
        outcome.messages.append(
            f"[Foci] Wrote {len(foci_table)} in-cell focus row(s) to {analysis_path.name} "
            f"(Foci sheet + Per-Cell columns).\n"
        )
    elif channel_rows:
        outcome.performed = True
        outcome.messages.append(
            f"[Foci] Assigned peaks but none fell inside cells for {mask_path.name}; "
            f"counts on Per-Cell Data are zero.\n"
        )
    else:
        outcome.messages.append(f"[Foci] No in-cell foci to write for {mask_path.name}.\n")

    return outcome


def manifest_to_run_info_dataframe(manifest: dict[str, Any]) -> pd.DataFrame:
    """Two-column Setting / Value table for the Run Info sheet."""
    rows: list[dict[str, str]] = []
    for key, value in manifest.items():
        if value is None or value == "":
            continue
        if isinstance(value, (list, tuple)):
            value = "; ".join(str(v) for v in value)
        rows.append({"Setting": str(key), "Value": str(value)})
    if not rows:
        rows.append({"Setting": "Note", "Value": "No run metadata recorded."})
    return pd.DataFrame(rows)


def write_run_info_sheet(analysis_path: Path, manifest: dict[str, Any]) -> None:
    """Insert or replace **Run Info** as the first sheet in the workbook."""
    from openpyxl import load_workbook

    analysis_path = Path(analysis_path)
    if not analysis_path.is_file():
        return

    df = manifest_to_run_info_dataframe(manifest)
    with pd.ExcelWriter(analysis_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        df.to_excel(writer, sheet_name=RUN_INFO_SHEET, index=False)

    try:
        wb = load_workbook(analysis_path)
        if RUN_INFO_SHEET in wb.sheetnames:
            idx = wb.sheetnames.index(RUN_INFO_SHEET)
            if idx > 0:
                wb.move_sheet(wb[RUN_INFO_SHEET], offset=-idx)
        wb.save(analysis_path)
        wb.close()
    except Exception:
        pass

    try:
        from bactomate.modules.excel_autofit import autofit_openpyxl_workbook
        from openpyxl import load_workbook

        wb = load_workbook(analysis_path)
        autofit_openpyxl_workbook(wb)
        wb.save(analysis_path)
        wb.close()
    except Exception:
        pass


def default_manifest_timestamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _sample_label_from_workbook_path(workbook_path: Path) -> str:
    stem = workbook_path.stem.replace("_analysis", "").replace("_masks_analysis", "")
    return stem


def _is_aggregate_analysis_workbook(path: Path) -> bool:
    name = path.name.lower()
    if name == MASTER_PER_CELL_SUMMARY_FILENAME.lower():
        return True
    if "combined_summary" in name:
        return True
    return False


# Generic fallback channel label when parsing fails (becomes ``foci_foci_*`` columns).
_FOCI_GENERIC_CHANNEL_PREFIX = "foci"
_FOCI_KNOWN_CHANNEL_PREFIXES = frozenset(
    {
        "gfp",
        "mcherry",
        "rfp",
        "cfp",
        "yfp",
        "dsred",
        "tdtomato",
        "mng",
        "scarlet",
        "mscarlet",
        "cy5",
        "cy3",
        "dapi",
    }
)

_FOCI_PER_CELL_METRIC_SUFFIXES = (
    "foci_count",
    "foci_intensities",
    "foci_mean_intensities",
    "foci_sizes_px",
    "foci_mean_intensity",
    "foci_total_intensity",
    "foci_max_intensity",
)


def _foci_channel_prefixes_from_columns(columns: Iterable[str]) -> set[str]:
    """Channel keys from ``{ch}_foci_count`` / ``foci_ids_{ch}`` Per-Cell columns."""
    prefixes: set[str] = set()
    for col in columns:
        name = str(col)
        if name.endswith("_foci_count"):
            prefixes.add(name[: -len("_foci_count")])
        elif name.startswith("foci_ids_"):
            prefixes.add(name[len("foci_ids_") :])
    return prefixes


def _preferred_foci_channel_prefix(prefixes: set[str]) -> str:
    """Pick a single fluorophore prefix when merging duplicate foci column families."""
    if not prefixes:
        return _FOCI_GENERIC_CHANNEL_PREFIX
    lower_map = {p.lower(): p for p in prefixes}
    if "gfp" in lower_map:
        return lower_map["gfp"]
    for key in sorted(lower_map.keys()):
        if key in _FOCI_KNOWN_CHANNEL_PREFIXES and key != _FOCI_GENERIC_CHANNEL_PREFIX:
            return lower_map[key]
    without_generic = [p for p in prefixes if p != _FOCI_GENERIC_CHANNEL_PREFIX]
    if without_generic:
        return sorted(without_generic, key=str.lower)[0]
    return "GFP"


def _canonicalize_per_cell_foci_columns(per_cell: pd.DataFrame) -> pd.DataFrame:
    """
    Merge duplicate foci column families in stacked Per-Cell exports.

    Older workbooks used the generic fallback label ``foci`` (``foci_foci_count``) even when
  the foci stack was ``*_foci_GFP``. Newer runs use ``GFP_foci_count``. When building the
    master summary, coalesce those into one channel column set.
    """
    if per_cell is None or per_cell.empty:
        return per_cell

    prefixes = _foci_channel_prefixes_from_columns(per_cell.columns)
    generic = _FOCI_GENERIC_CHANNEL_PREFIX
    if generic not in prefixes:
        return per_cell

    out = per_cell.copy()
    target = _preferred_foci_channel_prefix(prefixes)

    def _pair_cols(ch: str, suffix: str) -> tuple[str, str]:
        if suffix == "foci_count":
            return f"{ch}_foci_count", f"foci_ids_{ch}"
        return f"{ch}_{suffix}", f"foci_ids_{ch}"

    for suffix in _FOCI_PER_CELL_METRIC_SUFFIXES:
        src = f"{generic}_{suffix}"
        dst = f"{target}_{suffix}"
        if src not in out.columns:
            continue
        if dst not in out.columns:
            out[dst] = out[src]
        else:
            out[dst] = out[dst].combine_first(out[src])
        out.drop(columns=[src], inplace=True, errors="ignore")

    src_ids = f"foci_ids_{generic}"
    dst_ids = f"foci_ids_{target}"
    if src_ids in out.columns:
        if dst_ids not in out.columns:
            out[dst_ids] = out[src_ids]
        else:
            out[dst_ids] = out[dst_ids].combine_first(out[src_ids])
        out.drop(columns=[src_ids], inplace=True, errors="ignore")

    return out


def discover_analysis_workbooks(analysis_dir: Path) -> list[Path]:
    """List per-sample analysis workbooks in ``_analysis/`` (excludes master/combined aggregates)."""
    analysis_dir = Path(analysis_dir)
    if not analysis_dir.is_dir():
        return []
    found: list[Path] = []
    for pattern in ("*_analysis.xlsx", "*_masks_analysis.xlsx"):
        found.extend(analysis_dir.glob(pattern))
    unique: dict[str, Path] = {}
    for p in sorted(found, key=lambda x: x.name.lower()):
        if not p.is_file() or _is_aggregate_analysis_workbook(p):
            continue
        try:
            key = str(p.resolve())
        except Exception:
            key = str(p)
        unique[key] = p
    return list(unique.values())


def build_master_per_cell_dataframe(
    excel_paths: Iterable[Path],
) -> tuple[pd.DataFrame | None, list[str]]:
    """
    Stack **Per-Cell Data** sheets from multiple analysis workbooks (one row block per file).

    Adds :data:`MASTER_PER_CELL_SOURCE_COLUMN` (workbook filename) and
    :data:`MASTER_PER_CELL_SAMPLE_COLUMN` (short sample label).
    """
    frames: list[pd.DataFrame] = []
    warnings: list[str] = []
    for p in excel_paths:
        p = Path(p)
        if not p.is_file() or _is_aggregate_analysis_workbook(p):
            continue
        try:
            with pd.ExcelFile(p) as xl:
                if PER_CELL_DATA_SHEET not in xl.sheet_names:
                    warnings.append(f"{p.name}: no '{PER_CELL_DATA_SHEET}' sheet")
                    continue
                per_cell = pd.read_excel(xl, sheet_name=PER_CELL_DATA_SHEET)
        except Exception as exc:
            warnings.append(f"{p.name}: {exc}")
            continue
        if per_cell is None or per_cell.empty:
            warnings.append(f"{p.name}: empty '{PER_CELL_DATA_SHEET}' sheet")
            continue
        block = per_cell.copy()
        block = _canonicalize_per_cell_foci_columns(block)
        block.insert(0, MASTER_PER_CELL_SAMPLE_COLUMN, _sample_label_from_workbook_path(p))
        block.insert(0, MASTER_PER_CELL_SOURCE_COLUMN, p.name)
        frames.append(block)
    if not frames:
        return None, warnings
    long_df = pd.concat(frames, ignore_index=True, sort=False)
    long_df = _canonicalize_per_cell_foci_columns(long_df)
    lead = [MASTER_PER_CELL_SOURCE_COLUMN, MASTER_PER_CELL_SAMPLE_COLUMN]
    other = [c for c in long_df.columns if c not in lead]
    long_df = long_df[lead + other]
    return long_df, warnings


def write_master_per_cell_summary(
    analysis_dir: Path,
    *,
    excel_paths: Iterable[Path] | None = None,
    log_callback: Callable[[str], None] | None = None,
) -> bool:
    """
    Write ``@summary_per-cell-data.xlsx`` with all per-cell rows stacked on one sheet.

    When ``excel_paths`` is omitted, every analysis workbook in ``analysis_dir`` is included.
    Removes a stale master file when no per-cell rows are available.
    """
    analysis_dir = Path(analysis_dir)
    paths = list(excel_paths) if excel_paths is not None else discover_analysis_workbooks(analysis_dir)
    out_path = analysis_dir / MASTER_PER_CELL_SUMMARY_FILENAME

    def _log(msg: str) -> None:
        if log_callback is not None:
            log_callback(msg if msg.endswith("\n") else msg + "\n")

    long_df, warnings = build_master_per_cell_dataframe(paths)
    if long_df is None or long_df.empty:
        if out_path.is_file():
            try:
                out_path.unlink()
            except Exception:
                pass
        if warnings:
            _log(
                f"[Analyze] Master per-cell summary not written "
                f"({len(warnings)} workbook(s) had no per-cell data).\n"
            )
        return False

    analysis_dir.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        long_df.to_excel(writer, sheet_name=MASTER_PER_CELL_SUMMARY_SHEET, index=False)

    try:
        from bactomate.modules.excel_autofit import autofit_openpyxl_workbook
        from openpyxl import load_workbook

        wb = load_workbook(out_path)
        autofit_openpyxl_workbook(wb)
        wb.save(out_path)
        wb.close()
    except Exception:
        pass

    n_files = int(long_df[MASTER_PER_CELL_SOURCE_COLUMN].nunique())
    _log(
        f"[Analyze] Master per-cell summary saved: {out_path.name} "
        f"({len(long_df)} row(s) from {n_files} workbook(s)).\n"
    )
    if warnings:
        preview = ", ".join(warnings[:3])
        tail = f" (+{len(warnings) - 3} more)" if len(warnings) > 3 else ""
        _log(f"[Analyze] Master per-cell summary skipped: {preview}{tail}\n")
    return True
