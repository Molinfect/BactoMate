from __future__ import annotations
import re
import subprocess
import sys
from pathlib import Path
from typing import Any, Optional
import numpy as np
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication, QButtonGroup, QCheckBox, QDoubleSpinBox, QFrame, QGroupBox, QGridLayout, QHBoxLayout, QLabel, QLineEdit, QPushButton, QRadioButton, QSpinBox, QTextEdit, QVBoxLayout, QWidget
from bactomate.gui.analysis_paths import analysis_dir_from_base, base_dir_for_mask, stem_from_mask_path
from bactomate.gui.constants import PROJECT_ROOT
from bactomate.gui.preferences import get_default_browse_dir
from bactomate.gui.session_params import apply_to_widgets
from bactomate.gui.layout_helpers import style_compact_action_button, style_path_row_browse_button, vbox_from_rows
from bactomate.gui.styles import FLAT_GROUPBOX_FRAME_STYLESHEET
from bactomate.gui.widgets import create_info_icon, show_file_selection_dialog, show_linked_mask_selection_dialog, show_message_dialog
from bactomate.gui.workers import WorkerThread
from bactomate.gui.log_config import get_logger
from bactomate.modules.validation import mask_fi_spatial_shape_match, paths_exist
from bactomate.gui.tabs._analyze_masks_base import AnalyzeMasksBase

class AnalyzeTab(AnalyzeMasksBase):
    """Section 3: Analyze (cell/FI analysis + optional colocalization reporting; writes *_analysis*.xlsx)."""

    def _stage_key(self) -> str:
        return 'analyze'

    def _file_group_title(self) -> str:
        return 'Analyze artifacts'

    def _preview_button_label(self) -> str:
        return 'Preview'

    def _primary_button_label(self) -> str:
        return 'Analyze File/Folder'

    def _primary_button_tooltip(self) -> str:
        return 'Run analysis and write *_analysis.xlsx outputs'

    def _uses_combined_path_actions_group(self) -> bool:
        return True

    def _show_preview_action_in_combined_group(self) -> bool:
        return False

    def _extend_combined_path_actions_group(self, outer: QVBoxLayout, parent: QWidget) -> None:
        self._add_use_filtered_masks_row(outer, parent)

    def _wire_stage_buttons(self) -> None:
        try:
            self.preview_action_btn.setVisible(False)
        except Exception:
            pass
        self._analysis_in_progress = False
        self._cancel_analysis_requested = False
        self._analysis_worker: WorkerThread | None = None
        self.primary_action_btn.clicked.connect(self._on_analyze_or_cancel_clicked)

    def _on_analyze_or_cancel_clicked(self) -> None:
        if bool(getattr(self, '_analysis_in_progress', False)):
            self._cancel_analysis_requested = True
            w = getattr(self, '_analysis_worker', None)
            if w is not None and w.isRunning():
                try:
                    w.terminate_process()
                except Exception:
                    pass
            try:
                self.primary_action_btn.setEnabled(False)
                self.primary_action_btn.setText('Cancelling…')
            except Exception:
                pass
            return
        self.analyze_masks()

    def _log_analysis_run_params(self, *, selected_files: list[Path] | None=None) -> None:
        params = dict(self.get_session_params())
        params['selected_file_count'] = len(selected_files or [])
        params['selected_files'] = [Path(p).name for p in selected_files or []]
        lines = ['Analyze Masks run parameters:']
        for (key, value) in sorted(params.items()):
            lines.append(f'  {key}: {value}')
        if hasattr(self.parent_window, 'log_output'):
            self.parent_window.log_output('\n'.join(lines) + '\n')
        if getattr(self.parent_window, '_save_session_log', None):
            self.parent_window._save_session_log()

    def _reset_analyze_btn_idle(self) -> None:
        btn = self.primary_action_btn
        if hasattr(btn, 'set_progress_active'):
            btn.set_progress_active(False)
        btn.setText(self._primary_button_label())
        btn.setToolTip(self._primary_button_tooltip())
        btn.setEnabled(True)

    def _set_analyze_btn_cancel_mode(self, progress: int=0) -> None:
        btn = self.primary_action_btn
        if hasattr(btn, 'set_progress_active'):
            btn.set_progress_active(True)
            btn.set_progress(max(3, min(100, int(progress))))
        btn.setText('Cancel Analysis')
        btn.setToolTip('Stop the current analysis job.')
        btn.setEnabled(True)

    def _set_analysis_percent_ui(self, pct: int, *, label: str | None=None) -> None:
        pct = max(0, min(100, int(pct)))
        total = max(1, int(getattr(self, '_analysis_progress_total', 1)))
        try:
            pw = self.parent_window
            if hasattr(pw, 'progress_bar'):
                pw.progress_bar.setVisible(True)
                pw.progress_bar.setRange(0, 100)
                pw.progress_bar.setValue(pct)
            if label and hasattr(pw, 'set_status_message'):
                pw.set_status_message(label)
        except Exception:
            pass
        if hasattr(self.primary_action_btn, 'set_progress'):
            btn_pct = max(3, pct) if bool(getattr(self, '_analysis_in_progress', False)) and pct < 100 else pct
            self.primary_action_btn.set_progress(btn_pct)

    def _set_analysis_progress_ui(self, completed: float, *, label: str | None=None) -> None:
        total = max(1, int(getattr(self, '_analysis_progress_total', 1)))
        completed_f = max(0.0, min(float(completed), float(total)))
        pct = int(min(100, round(100.0 * completed_f / float(total))))
        self._set_analysis_percent_ui(pct, label=label)

    def _start_analysis_progress(self, total: int, label: str='Analyzing masks...') -> None:
        self._analysis_progress_total = max(1, int(total))
        self._set_analyze_btn_cancel_mode(0)
        self._set_analysis_progress_ui(0, label=label)
        QApplication.processEvents()

    def _update_analysis_progress(self, value: int, label: str | None=None) -> None:
        self._set_analysis_progress_ui(value, label=label)
        QApplication.processEvents()

    def _close_analysis_progress(self, final_value: int | None=None) -> None:
        if final_value is not None:
            self._set_analysis_progress_ui(int(final_value))
        try:
            pw = self.parent_window
            if hasattr(pw, 'progress_bar'):
                pw.progress_bar.setVisible(False)
                pw.progress_bar.setValue(0)
            if hasattr(pw, 'set_status_message'):
                pw.set_status_message('')
        except Exception:
            pass

    def _set_post_analysis_progress(self, label: str, *, percent: int=90) -> None:
        """Keep the in-button progress visible during post-analysis (foci, vis, Run Info)."""
        try:
            self._set_analyze_btn_cancel_mode(int(percent))
            self._set_analysis_percent_ui(int(percent), label=label)
            QApplication.processEvents()
        except Exception:
            pass

    def _handle_analysis_worker_progress(self, message: str) -> None:
        text = str(message or '')
        if hasattr(self.parent_window, 'log_output') and text:
            self.parent_window.log_output(text)
        match = re.search('\\[ANALYZE_PROGRESS\\]\\s+phase=(\\w+)\\s+current=(\\d+)\\s+total=(\\d+)', text)
        if not match:
            return
        (phase, current_s, total_s) = match.groups()
        try:
            current = max(0, int(current_s))
            total = max(1, int(total_s))
        except ValueError:
            return
        frac = max(0.0, min(1.0, current / float(total)))
        idx = int(getattr(self, 'current_task_index', 0) or 0)
        task_completed = idx + 0.85 * frac
        label = f'Analyzing frames ({current}/{total})...'
        if phase != 'frames':
            label = f'Analyzing {phase} ({current}/{total})...'
        self._set_analysis_progress_ui(task_completed, label=label)
        QApplication.processEvents()

    def _analysis_output_folder_from_paths(self, paths: list[str | Path] | None) -> Path | None:
        try:
            output_dirs = {Path(p).parent.resolve() for p in paths or [] if p}
            if len(output_dirs) == 1:
                return next(iter(output_dirs))
            if output_dirs:
                selected = Path(self.file_path.text())
                return selected if selected.is_dir() else selected.parent
        except Exception:
            pass
        return None

    def _add_use_filtered_masks_row(self, outer: QVBoxLayout, parent: QWidget) -> None:
        self.use_filtered_masks_cb = QCheckBox('Use filtered masks if available')
        self.use_filtered_masks_cb.setChecked(True)
        self.use_filtered_masks_cb.setToolTip('When checked, analysis will prefer filtered masks in\n`filtered_masks/*_mask_filtered.tif` under each segmentation folder.\nIf none are found, it falls back to `*_masks.tif`.')
        self.filtered_masks_hint = QLabel('')
        self.filtered_masks_hint.setStyleSheet('color: #9e9e9e; font-size: 9pt;')
        self.filtered_masks_hint.setWordWrap(True)
        filter_row = QHBoxLayout()
        filter_row.setSpacing(8)
        filter_row.addWidget(self.use_filtered_masks_cb, 0)
        filter_row.addWidget(self.filtered_masks_hint, 1)
        outer.addLayout(filter_row)

        def _refresh_hint() -> None:
            try:
                txt = (self.file_path.text() or '').strip()
                if not txt:
                    self.filtered_masks_hint.setText('')
                    return
                p = Path(txt)
                if not p.exists() or not self.use_filtered_masks_cb.isChecked():
                    self.filtered_masks_hint.setText('')
                    return
                folder = p if p.is_dir() else p.parent
                from bactomate.gui.artifact_paths import collect_filtered_mask_files_in_folder
                n_filt = len(collect_filtered_mask_files_in_folder(p))
                if n_filt:
                    self.filtered_masks_hint.setText(f'({n_filt} filtered mask(s) found)')
                else:
                    self.filtered_masks_hint.setText('(no filtered masks found; will fall back)')
            except Exception:
                self.filtered_masks_hint.setText('')
        self.file_path.textChanged.connect(lambda : _refresh_hint())
        self.use_filtered_masks_cb.toggled.connect(lambda _v: _refresh_hint())
        _refresh_hint()

    def _mask_paths_for_foci_csv_hint(self, path: Path) -> list[Path]:
        """Mask files used to estimate how many foci CSVs are available for the current path."""
        path = Path(path)
        if not path.exists():
            return []
        if path.is_file():
            return [path.resolve()]
        folder = path.resolve()
        mask_set: set[Path] = set()
        use_filtered = bool(getattr(self, 'use_filtered_masks_cb', None) and self.use_filtered_masks_cb.isChecked())
        if use_filtered:
            from bactomate.gui.artifact_paths import collect_filtered_mask_files_in_folder
            mask_set.update(collect_filtered_mask_files_in_folder(folder))
        if not mask_set:
            for pat in ('*_masks.tif', '*_masks.tiff', '*_mask_filtered.tif', '*_mask_filtered.tiff'):
                mask_set.update(folder.glob(pat))
            for seg in sorted(folder.glob('_segmentation*')):
                if seg.is_dir():
                    for pat in ('*_masks.tif', '*_masks.tiff', '*_mask_filtered.tif', '*_mask_filtered.tiff'):
                        mask_set.update(seg.glob(pat))
                    filt_dir = seg / 'filtered_masks'
                    if filt_dir.is_dir():
                        mask_set.update(filt_dir.glob('*_mask_filtered.tif'))
                        mask_set.update(filt_dir.glob('*_mask_filtered.tiff'))
        return sorted(mask_set, key=lambda p: p.name.lower())

    def _count_unique_foci_csv_artifacts(self, mask_paths: list[Path], *, prefer_filtered: bool) -> tuple[int, set[str]]:
        """Return (count, csv_folder names) of unique importable foci CSVs."""
        from bactomate.gui.analysis_workbook import discover_foci_csv_artifacts_for_mask
        seen_csv: set[str] = set()
        folders: set[str] = set()
        for mp in mask_paths:
            try:
                artifacts = discover_foci_csv_artifacts_for_mask(self, mp, prefer_filtered=prefer_filtered)
            except Exception:
                continue
            for art in artifacts:
                key = str(art.csv_path.resolve())
                if key in seen_csv:
                    continue
                seen_csv.add(key)
                folders.add(art.csv_folder)
        return (len(seen_csv), folders)

    def _foci_csv_counts_for_masks(self, mask_paths: list[Path]) -> tuple[int, int]:
        """Return (n filtered foci CSV, n detection-only CSV) across masks/channels."""
        from bactomate.gui.artifact_paths import infer_pipeline_dir_from_mask, resolve_filtered_foci_csv, resolve_segmentation_foci_csv
        seen_filtered: set[str] = set()
        seen_detection: set[str] = set()
        for mp in mask_paths:
            try:
                pipeline_dir = infer_pipeline_dir_from_mask(Path(mp))
                sources = self._resolve_foci_paths_for_mask(Path(mp))
            except Exception:
                continue
            for src in sources or []:
                try:
                    filt = resolve_filtered_foci_csv(pipeline_dir, mask_path=Path(mp), foci_source_path=Path(src))
                except Exception:
                    filt = None
                if filt is not None and Path(filt).is_file():
                    key = str(Path(filt).resolve())
                    if key not in seen_filtered:
                        seen_filtered.add(key)
                    continue
                try:
                    det = resolve_segmentation_foci_csv(pipeline_dir, mask_path=Path(mp), foci_source_path=Path(src))
                except Exception:
                    det = None
                if det is not None and Path(det).is_file():
                    key = str(Path(det).resolve())
                    if key not in seen_detection:
                        seen_detection.add(key)
        return (len(seen_filtered), len(seen_detection))

    def _fi_tiff_count_for_masks(self, mask_paths: list[Path]) -> int:
        """Unique FI-channel TIFFs (*_fi*, *_fi_foci_*, etc.) name-matched to masks."""
        seen: set[str] = set()
        for mp in mask_paths:
            try:
                for p in self._resolve_fi_paths_for_mask(Path(mp)) or []:
                    key = str(Path(p).resolve())
                    seen.add(key)
            except Exception:
                continue
        return len(seen)

    def _companions_for_analysis_mask(self, mask_path: Path) -> dict[str, list]:
        """Foci CSV and FI paths linked to one mask for the analysis file picker."""
        from bactomate.gui.artifact_paths import infer_pipeline_dir_from_mask, resolve_filtered_foci_csv, resolve_segmentation_foci_csv
        mask_path = Path(mask_path)
        include_foci_csv = bool(getattr(self, 'include_foci_from_detection', None) and self.include_foci_from_detection.isChecked())
        foci_rows: list[tuple[Path, str]] = []
        seen_csv: set[str] = set()
        if include_foci_csv:
            try:
                pipeline_dir = infer_pipeline_dir_from_mask(mask_path)
                for src in self._resolve_foci_paths_for_mask(mask_path) or []:
                    src_p = Path(src)
                    filt = resolve_filtered_foci_csv(pipeline_dir, mask_path=mask_path, foci_source_path=src_p)
                    if filt is not None and Path(filt).is_file():
                        key = str(Path(filt).resolve())
                        if key not in seen_csv:
                            seen_csv.add(key)
                            foci_rows.append((Path(filt), 'filtered'))
                        continue
                    det = resolve_segmentation_foci_csv(pipeline_dir, mask_path=mask_path, foci_source_path=src_p)
                    if det is not None and Path(det).is_file():
                        key = str(Path(det).resolve())
                        if key not in seen_csv:
                            seen_csv.add(key)
                            foci_rows.append((Path(det), 'detection'))
            except Exception:
                pass
        fi_rows: list[tuple[Path, str]] = []
        try:
            if self._is_match_by_name_mode():
                explicit = [Path(p) for p in self._resolve_fi_paths_for_mask(mask_path) or []]
                explicit_keys = {str(p.resolve()) for p in explicit}
                for p in explicit:
                    fi_rows.append((p, 'fi'))
                for p in self._resolve_foci_paths_for_mask(mask_path) or []:
                    pp = Path(p)
                    key = str(pp.resolve())
                    if key in explicit_keys:
                        continue
                    fi_rows.append((pp, 'fi foci channel'))
            else:
                single = getattr(self, '_single_fi_path', None)
                if single is not None and Path(single).is_file():
                    fi_rows.append((Path(single), 'fi'))
        except Exception:
            pass
        return {'foci': foci_rows, 'fi': fi_rows}

    def _mask_has_analysis_companions(self, mask_path: Path) -> bool:
        comp = self._companions_for_analysis_mask(mask_path)
        return bool(comp.get('foci') or comp.get('fi'))

    def _refresh_foci_csv_hint(self) -> None:
        """Update foci CSV hint/checkbox and auto-enable Analyze FI when *_fi* TIFFs exist."""
        hint = getattr(self, 'foci_csv_hint', None)
        cb = getattr(self, 'include_foci_from_detection', None)
        fi_cb = getattr(self, 'analyze_fi', None)

        def _set_fi_checked(on: bool) -> None:
            if fi_cb is None:
                return
            fi_cb.blockSignals(True)
            fi_cb.setChecked(on)
            fi_cb.blockSignals(False)
            try:
                self._update_fi_ui_state()
            except Exception:
                pass
        if hint is None:
            return
        try:
            txt = (self.file_path.text() or '').strip()
            if not txt:
                hint.setText('')
                if cb is not None:
                    cb.blockSignals(True)
                    cb.setChecked(False)
                    cb.blockSignals(False)
                _set_fi_checked(False)
                return
            path = Path(txt)
            if not path.exists():
                hint.setText('')
                if cb is not None:
                    cb.blockSignals(True)
                    cb.setChecked(False)
                    cb.blockSignals(False)
                _set_fi_checked(False)
                return
            masks = self._mask_paths_for_foci_csv_hint(path)
            if not masks:
                hint.setText('(no masks found for path)')
                if cb is not None:
                    cb.blockSignals(True)
                    cb.setChecked(False)
                    cb.blockSignals(False)
                _set_fi_checked(False)
                return
            n_fi = self._fi_tiff_count_for_masks(masks)
            _set_fi_checked(n_fi > 0)
            (n_filtered, n_detection) = self._foci_csv_counts_for_masks(masks)
            has_any = n_filtered > 0 or n_detection > 0
            if cb is not None:
                cb.blockSignals(True)
                cb.setChecked(has_any)
                cb.blockSignals(False)
            if n_filtered:
                fl = 'file' if n_filtered == 1 else 'files'
                if n_detection:
                    dl = 'file' if n_detection == 1 else 'files'
                    hint.setText(f'({n_filtered} filtered foci CSV {fl} found — preferred; {n_detection} segmentation detection {dl} where filtered is missing)')
                else:
                    hint.setText(f'({n_filtered} filtered foci CSV {fl} found — preferred)')
            elif n_detection:
                dl = 'file' if n_detection == 1 else 'files'
                hint.setText(f'({n_detection} segmentation foci CSV {dl} found; no filtered foci CSV)')
            else:
                hint.setText('(no foci CSV found)')
        except Exception:
            hint.setText('')

    def _build_stage_analysis_subpanel(self) -> Optional[QWidget]:
        """
        Analyze controls for Section 3.

        Layout overview:
            ┌─ Fluorescence intensity (FI) ─────────────────────────────────┐
            │ [ ] Analyze FI [info]   Channels: ( ) One file (•) Match name │
            │ [Select FI file…]  (only for one-file mode)                   │
            │ [ ] Contour FI   [Preview Contour FI…]                        │
            │ ┌ feedback box (visible only while Analyze FI is on) ────────┐│
            │ └─────────────────────────────────────────────────────────────┘│
            └────────────────────────────────────────────────────────────────┘
            ┌─ Foci analysis (shared foci CSV) ───────────────────────────────┐
            │ [ ] Include foci from segmentation CSV [info] (N … found)   │
            │ ─────────────────────────────────────────────────────────── │
            │ [ ] Add colocalization sheets per channel pair       [info]   │
            │ Max distance / threshold / same cell only                   │
            └────────────────────────────────────────────────────────────────┘
        """
        outer = QWidget()
        outer.setMinimumWidth(0)
        outer_layout = QVBoxLayout(outer)
        outer_layout.setContentsMargins(0, 6, 0, 0)
        outer_layout.setSpacing(6)
        self._single_fi_path: Optional[Path] = None
        fi_group = QGroupBox('Fluorescence intensity (FI)')
        fi_group.setFlat(True)
        fi_group.setStyleSheet(FLAT_GROUPBOX_FRAME_STYLESHEET)
        fi_group.setMinimumWidth(0)
        fi_layout = QVBoxLayout()
        fi_layout.setSpacing(6)
        fi_layout.setContentsMargins(8, 6, 8, 6)
        self.analyze_fi = QCheckBox('Analyze FI')
        self.analyze_fi.setChecked(False)
        ch_label = QLabel('Channels:')
        ch_label.setStyleSheet('color: #9e9e9e;')
        self.fi_mode_one_file = QRadioButton('One file for all masks')
        self.fi_mode_match_by_name = QRadioButton('Match by name')
        self.fi_mode_match_by_name.setChecked(True)
        self.fi_mode_group = QButtonGroup(self)
        self.fi_mode_group.addButton(self.fi_mode_one_file)
        self.fi_mode_group.addButton(self.fi_mode_match_by_name)
        fi_row = QHBoxLayout()
        fi_row.setSpacing(8)
        fi_row.addWidget(self.analyze_fi, 0)
        fi_row.addWidget(create_info_icon('Analyze fluorescence intensity (FI) from a separate channel.\n\nMatch by name will search for `stem_fi.tif` / `stem_fi_*` next to the mask.', parent_widget=self), 0)
        fi_row.addSpacing(10)
        fi_row.addWidget(ch_label, 0)
        fi_row.addWidget(self.fi_mode_one_file, 0)
        fi_row.addWidget(self.fi_mode_match_by_name, 0)
        self.fi_select_btn = QPushButton('Select FI file…')
        style_path_row_browse_button(self.fi_select_btn)
        self.fi_select_btn.clicked.connect(self._select_fi_file)
        self.analyze_fi_contour = QCheckBox('Contour FI')
        self.analyze_fi_contour.setChecked(False)
        self.analyze_fi_contour.setToolTip('Compute FI on a thin contour ring around each cell instead of the full mask.')
        self.fi_contour_thickness = QSpinBox()
        self.fi_contour_thickness.setRange(1, 10)
        self.fi_contour_thickness.setValue(1)
        self.fi_contour_thickness.setSuffix(' px')
        self.fi_contour_thickness.hide()
        self.fi_contour_preview_btn = QPushButton('Preview Contour FI…')
        style_compact_action_button(self.fi_contour_preview_btn, horizontal_padding=10)
        self.fi_contour_preview_btn.setMinimumWidth(0)
        self.fi_contour_preview_btn.setMaximumWidth(170)
        self.fi_contour_preview_btn.setEnabled(False)
        self.fi_contour_preview_btn.setToolTip('Select a mask, then browse input/FI channels with Contour FI overlay and LUT.')
        self.fi_contour_preview_btn.clicked.connect(self._show_fi_contour_on_filter_preview)
        fi_row.addWidget(self.fi_select_btn, 0)
        fi_row.addStretch(1)
        fi_layout.addLayout(fi_row)
        contour_row = QHBoxLayout()
        contour_row.setSpacing(8)
        contour_row.addWidget(self.analyze_fi_contour, 0)
        contour_row.addWidget(self.fi_contour_thickness, 0)
        contour_row.addWidget(self.fi_contour_preview_btn, 0)
        contour_row.addStretch(1)
        fi_layout.addLayout(contour_row)
        self.fi_feedback = QTextEdit()
        self.fi_feedback.setReadOnly(True)
        self.fi_feedback.setMaximumHeight(110)
        self.fi_feedback.setPlaceholderText('Set mask path and choose FI mode to see what will be analyzed.')
        fi_layout.addWidget(self.fi_feedback)
        fi_group.setLayout(fi_layout)
        outer_layout.addWidget(fi_group)
        _foci_csv_shared = 'Uses Spotiflow peak CSV data from segmentation (Detect foci)—not Analyze FI.\n\nWhen “Use filtered masks if available” is on, each channel uses its foci_filtered CSV if present, otherwise foci_detection. Run Info lists the exact files used.\n\nRequires matching *_foci*.tif stacks. No re-detection in Analyze.'
        _foci_import_info = 'After analysis, import peak CSV data into the same *_analysis.xlsx.\n\nAdds a Foci sheet (foci_id ↔ Cell_ID), per-cell foci counts, foci_ids columns, peak intensities, and Summary foci totals.\n\n' + _foci_csv_shared
        _foci_coloc_info = 'Colocalization (foci channel pairs)\n\nWhen enabled, Analyze adds one Colocalization sheet per *_foci channel pair (e.g. Channel_1×Channel_2) to the same *_analysis.xlsx. Center-to-center distances are in pixels on the same image. You need at least two foci channels with peak CSV data. Independent from Analyze FI.\n\nHow matching works (three steps):\n1) Max distance — search radius only. For each focus in channel A, BactoMate lists candidate partners in channel B whose peak centers lie within this distance. This step does not mark foci as colocalized; it only decides who may be compared.\n2) One-to-one pairing. Among those candidates, BactoMate picks the best non-overlapping assignment (Hungarian algorithm). Each focus matches at most one partner in the other channel.\n3) Threshold — colocalized or not. A matched pair counts as colocalized only if its distance is ≤ Threshold (pixels). Pairs within Max distance but above Threshold still appear in the sheet with is_valid = False and are not counted in coloc_pairs_count on the Per-Cell sheet. coloc_match_index is one score per cell: coloc_pairs / max(channel foci counts) as %, or no foci / no correlation.\n\nExample: Max distance 8 px can find four candidate pairs in one cell; Threshold 1 px may count only pairs ≤ 1 px apart (e.g. 1.0 px yes, 1.4 px no).\n\nSame cell only:\n• On (default): Only pairs where both peaks share the same Cell_ID (same segmented bacterium). Best for two nuclear/cytoplasmic markers in the same cell.\n• Off: Matching is by image distance only — peaks may have different Cell_IDs. The sheet still lists cell_id_primary, cell_id_secondary, and same_cell (True/False). Use this when biology is not “same object, same cell”, but “are these two signals close in the image?”\nExample (off): A translational fusion reporter inside the rod (assigned to that cell’s Cell_ID) and extracellular phage or virus particles stained outside (often edge/background peaks, or another Cell_ID). You can ask whether reporter foci colocalize with nearby phage labels even though they are not in the same segmented region. Tune max distance/threshold for attachment distance.\n\nFoci assignment matches Filtering / “Include foci from segmentation CSV”: peaks on the cell edge (background pixel but within 2 px of the mask) are assigned to that cell. foci_id numbers restart at 1 in each channel.\n\n' + _foci_csv_shared
        foci_analysis_group = QGroupBox('Foci analysis')
        foci_analysis_group.setFlat(True)
        foci_analysis_group.setStyleSheet(FLAT_GROUPBOX_FRAME_STYLESHEET)
        foci_analysis_group.setMinimumWidth(0)
        foci_analysis_layout = QVBoxLayout()
        foci_analysis_layout.setSpacing(6)
        foci_analysis_layout.setContentsMargins(8, 6, 8, 6)
        self.include_foci_from_detection = QCheckBox('Include foci from segmentation CSV')
        self.include_foci_from_detection.setChecked(False)
        self.foci_csv_hint = QLabel('')
        self.foci_csv_hint.setStyleSheet('color: #9e9e9e; font-size: 9pt;')
        self.foci_csv_hint.setWordWrap(True)
        self.coloc_calc_all_pairs = QCheckBox('Colocalization sheets')
        self.coloc_calc_all_pairs.setChecked(False)
        self.coloc_calc_all_pairs.setToolTip('Add one colocalization sheet per *_foci channel pair.')
        foci_top = QHBoxLayout()
        foci_top.setSpacing(8)
        foci_top.addWidget(self.include_foci_from_detection, 0)
        foci_top.addWidget(create_info_icon(_foci_import_info, parent_widget=self), 0)
        foci_top.addSpacing(10)
        foci_top.addWidget(self.coloc_calc_all_pairs, 0)
        foci_top.addWidget(create_info_icon(_foci_coloc_info, parent_widget=self), 0)
        foci_top.addWidget(self.foci_csv_hint, 1)
        foci_analysis_layout.addLayout(foci_top)
        foci_analysis_rule = QFrame()
        foci_analysis_rule.setFrameShape(QFrame.Shape.HLine)
        foci_analysis_rule.setFrameShadow(QFrame.Shadow.Plain)
        foci_analysis_rule.setStyleSheet('QFrame { background-color: #3d4d62; border: none; max-height: 1px; margin-top: 2px; margin-bottom: 2px; }')
        foci_analysis_layout.addWidget(foci_analysis_rule)
        coloc_row = QHBoxLayout()
        coloc_row.setSpacing(8)
        self.coloc_merge_max_distance = QDoubleSpinBox()
        self.coloc_merge_max_distance.setRange(0.1, 500.0)
        self.coloc_merge_max_distance.setDecimals(1)
        self.coloc_merge_max_distance.setValue(8.0)
        self.coloc_merge_max_distance.setSuffix(' px')
        self.coloc_merge_max_distance.setMinimumWidth(72)
        self.coloc_merge_max_distance.setMaximumWidth(88)
        self.coloc_merge_threshold = QDoubleSpinBox()
        self.coloc_merge_threshold.setRange(0.1, 500.0)
        self.coloc_merge_threshold.setDecimals(1)
        self.coloc_merge_threshold.setValue(4.0)
        self.coloc_merge_threshold.setSuffix(' px')
        self.coloc_merge_threshold.setMinimumWidth(72)
        self.coloc_merge_threshold.setMaximumWidth(88)
        self.coloc_merge_same_cell_only = QCheckBox('Same cell only')
        self.coloc_merge_same_cell_only.setChecked(True)
        coloc_row.addSpacing(10)
        coloc_row.addWidget(QLabel('Max distance:'), 0)
        coloc_row.addWidget(self.coloc_merge_max_distance, 0)
        coloc_row.addWidget(QLabel('Threshold:'), 0)
        coloc_row.addWidget(self.coloc_merge_threshold, 0)
        coloc_row.addWidget(self.coloc_merge_same_cell_only, 0)
        coloc_row.addStretch(1)
        foci_analysis_layout.addLayout(coloc_row)
        self._coloc_param_widgets = (self.coloc_merge_max_distance, self.coloc_merge_threshold, self.coloc_merge_same_cell_only)
        foci_analysis_group.setLayout(foci_analysis_layout)
        outer_layout.addWidget(foci_analysis_group)
        self.analyze_fi.toggled.connect(self._update_fi_ui_state)
        self.fi_mode_one_file.toggled.connect(self._update_fi_ui_state)
        self.fi_mode_match_by_name.toggled.connect(self._update_fi_ui_state)
        self.analyze_fi_contour.toggled.connect(self._update_fi_contour_state)
        self.coloc_calc_all_pairs.toggled.connect(self._update_coloc_param_state)
        self.file_path.textChanged.connect(self._update_fi_feedback)
        self.file_path.textChanged.connect(lambda : self._refresh_foci_csv_hint())
        if getattr(self, 'use_filtered_masks_cb', None) is not None:
            self.use_filtered_masks_cb.toggled.connect(lambda _v: self._refresh_foci_csv_hint())
        self._update_fi_ui_state()
        self._update_coloc_param_state()
        self._refresh_foci_csv_hint()
        return outer

    def _update_coloc_param_state(self, *_args) -> None:
        """Grey out colocalization parameters when the feature is disabled."""
        on = bool(getattr(self, 'coloc_calc_all_pairs', None) and self.coloc_calc_all_pairs.isChecked())
        for w in getattr(self, '_coloc_param_widgets', ()) or ():
            try:
                w.setEnabled(on)
            except Exception:
                pass

    def _select_fi_file(self) -> None:
        from PyQt6.QtWidgets import QFileDialog
        (file_path, _) = QFileDialog.getOpenFileName(self, 'Select FI Channel TIFF File', get_default_browse_dir(), 'TIFF Files (*.tif *.tiff);;All Files (*)')
        if file_path:
            self._single_fi_path = Path(file_path)
            self._update_fi_feedback()

    def _is_match_by_name_mode(self) -> bool:
        return self._safe_checked(getattr(self, 'fi_mode_match_by_name', None), default=True)

    @staticmethod
    def _safe_checked(widget: Any, *, default: bool=False) -> bool:
        try:
            if widget is None or not hasattr(widget, 'isChecked'):
                return bool(default)
            return bool(widget.isChecked())
        except RuntimeError:
            return bool(default)

    @staticmethod
    def _safe_value(widget: Any, *, default: float=0.0) -> float:
        try:
            if widget is None or not hasattr(widget, 'value'):
                return float(default)
            return float(widget.value())
        except (RuntimeError, TypeError, ValueError):
            return float(default)

    @staticmethod
    def _safe_set_checked(widget: Any, checked: bool) -> None:
        try:
            if widget is not None and hasattr(widget, 'setChecked'):
                widget.setChecked(bool(checked))
        except RuntimeError:
            pass

    def _update_fi_contour_state(self, *_args) -> None:
        analyze_fi_on = self._safe_checked(getattr(self, 'analyze_fi', None), default=False)
        try:
            self.analyze_fi_contour.setEnabled(analyze_fi_on)
        except Exception:
            pass
        self._update_fi_contour_preview_button_state()

    def _update_fi_contour_preview_button_state(self) -> None:
        btn = getattr(self, 'fi_contour_preview_btn', None)
        if btn is None:
            return
        from bactomate.gui.fi_contour_filter_preview import fi_contour_preview_button_enabled
        (enabled, tip) = fi_contour_preview_button_enabled(self)
        try:
            btn.setEnabled(enabled)
            btn.setToolTip(tip)
        except Exception:
            pass

    def _show_fi_contour_on_filter_preview(self) -> None:
        from bactomate.gui.fi_contour_filter_preview import show_fi_contour_preview
        err = show_fi_contour_preview(self)
        if err:
            show_message_dialog(self, 'Contour FI preview', err, kind='warning')

    def _update_fi_ui_state(self, _checked=None) -> None:
        analyze_fi_on = self._safe_checked(getattr(self, 'analyze_fi', None), default=False)
        for w in (getattr(self, 'fi_mode_one_file', None), getattr(self, 'fi_mode_match_by_name', None)):
            if w is None:
                continue
            try:
                w.setEnabled(analyze_fi_on)
            except Exception:
                pass
        try:
            self.fi_select_btn.setVisible(analyze_fi_on and (not self._is_match_by_name_mode()))
            self.fi_select_btn.setEnabled(analyze_fi_on)
        except Exception:
            pass
        try:
            self.fi_feedback.setVisible(analyze_fi_on)
        except Exception:
            pass
        self._update_fi_contour_state()
        self._update_fi_feedback()

    def _update_fi_feedback(self) -> None:
        if not getattr(self, 'fi_feedback', None):
            return
        if not self._safe_checked(getattr(self, 'analyze_fi', None), default=False):
            self.fi_feedback.setPlainText('Analyze FI is off. Enable it to see FI pairings.\nColocalization remains available and runs from existing foci CSV artifacts.')
            return
        if self._is_match_by_name_mode():
            masks = self._get_mask_files_for_feedback()
            if not masks:
                self.fi_feedback.setPlainText('Set mask path (file or folder) to see FI pairings.')
                return
            fi_lines = []
            total = 0
            for mask_path in masks:
                fi_paths = self._resolve_fi_inputs_for_mask(mask_path)
                total += len(fi_paths)
                fi_lines.append(f"• {mask_path.name} → {(', '.join((p.name for p in fi_paths)) if fi_paths else '(no FI found)')}")
            header = f'{len(masks)} mask file(s) → {total} analyses\n\n'
            self.fi_feedback.setPlainText(header + 'Analyze FI (mask → FI channel(s)):\n' + '\n'.join(fi_lines))
        elif self._single_fi_path and self._single_fi_path.exists():
            self.fi_feedback.setPlainText(f'FI file: {self._single_fi_path.name} will be used for all masks.')
        else:
            self.fi_feedback.setPlainText('Select one FI file (button above) to use for all masks.')

    def analyze_masks(self) -> None:
        if bool(getattr(self, '_analysis_in_progress', False)):
            return
        if not self.file_path.text():
            show_message_dialog(self, 'No Selection', 'Please select a mask file or folder first.', kind='warning')
            return
        if not Path(self.file_path.text()).exists():
            show_message_dialog(self, 'Not Found', 'The selected file or folder does not exist.', kind='warning')
            return
        self._analyze_script_module = 'bactomate.scripts.analyze_cell_areas'
        self._post_analysis_coloc_masks: list[Path] | None = None
        is_folder = self._is_folder_mode()
        if is_folder:
            self._folder_failures: list[tuple[str, str]] = []
            mask_files = self._get_mask_files_for_feedback()
            if not mask_files:
                show_message_dialog(self, 'No Mask Files Found', 'No mask files found in the selected folder.', kind='warning')
                return
            include_foci_csv = bool(getattr(self, 'include_foci_from_detection', None) and self.include_foci_from_detection.isChecked())
            show_fi_picker = bool(self.analyze_fi.isChecked())
            use_linked_picker = any((self._mask_has_analysis_companions(m) for m in mask_files))
            if use_linked_picker and include_foci_csv and show_fi_picker:
                picker_message = f'Select mask files to analyze ({len(mask_files)} listed). Linked foci CSV and FI TIFFs follow each mask on the right.'
                right_header = 'Linked foci CSV + FI-channel TIFFs'
            elif use_linked_picker and include_foci_csv:
                picker_message = f'Select mask files to analyze ({len(mask_files)} listed). Linked foci CSV files follow each mask on the right.'
                right_header = 'Linked foci CSV files'
            elif use_linked_picker and show_fi_picker:
                picker_message = f'Select mask files to analyze ({len(mask_files)} listed). Linked FI-channel TIFFs follow each mask on the right.'
                right_header = 'Linked FI-channel TIFFs'
            else:
                picker_message = f'Select mask files to analyze (all {len(mask_files)} files selected by default):'
                right_header = 'Linked artifacts'
            if use_linked_picker:
                selected_files = show_linked_mask_selection_dialog(self, mask_files, companion_for_mask=self._companions_for_analysis_mask, title='Select Mask Files for Analysis', message=picker_message, right_header=right_header, show_fi=show_fi_picker)
            else:
                selected_files = show_file_selection_dialog(self, mask_files, title='Select Mask Files for Analysis', message=picker_message)
            if selected_files is None or len(selected_files) == 0:
                return
            if self.analyze_fi.isChecked() and (not self._is_match_by_name_mode()) and (not (getattr(self, '_single_fi_path', None) and self._single_fi_path.exists())):
                show_message_dialog(self, 'FI File Required', "Analyze FI is on with 'One FI file for all masks' but no file is selected.", kind='warning')
                return
            self._analysis_in_progress = True
            self._cancel_analysis_requested = False
            self._set_analyze_btn_cancel_mode(0)
            self.folder_files = selected_files
            self.folder_results = []
            self.folder_output_paths = []
            batch_fi_on = self.analyze_fi.isChecked() and self._is_match_by_name_mode()
            single_fi_path = getattr(self, '_single_fi_path', None) if self.analyze_fi.isChecked() and (not batch_fi_on) else None
            if single_fi_path and (not single_fi_path.exists()):
                single_fi_path = None
            self.folder_fi_tasks: list[tuple[Path, Path | None, Path]] = []
            self.folder_phase_per_mask: dict[Path, Path | None] = {}

            def _add_task(mask_p: Path, fi_p: Path | None, out_p: Path) -> None:
                self.folder_fi_tasks.append((mask_p, fi_p, out_p))
            for mask_path in selected_files:
                mask_path = mask_path.resolve()
                base_dir = base_dir_for_mask(mask_path)
                stem = stem_from_mask_path(mask_path)
                analysis_dir = analysis_dir_from_base(base_dir)
                if batch_fi_on:
                    fi_paths = self._resolve_fi_inputs_for_mask(mask_path)
                    if not fi_paths:
                        _add_task(mask_path, None, analysis_dir / f'{stem}_analysis.xlsx')
                    else:
                        for fi_path in fi_paths:
                            suffix = self._fi_output_suffix(fi_path, stem)
                            _add_task(mask_path, fi_path, analysis_dir / f'{stem}_analysis{suffix}.xlsx')
                else:
                    _add_task(mask_path, single_fi_path, analysis_dir / f'{stem}_analysis.xlsx')
            validation_errors = []
            for (mask_path, fi_path, _out) in self.folder_fi_tasks:
                (ok, _missing) = paths_exist([mask_path])
                if not ok:
                    validation_errors.append(f'{mask_path.name}: mask file not found.')
                    continue
                if fi_path is not None:
                    (ok, _missing) = paths_exist([fi_path])
                    if not ok:
                        validation_errors.append(f'{mask_path.name}: FI file not found.')
                        continue
                    (ok, err) = mask_fi_spatial_shape_match(mask_path, fi_path)
                    if not ok and err:
                        validation_errors.append(f'{mask_path.name}: {err}')
            if validation_errors:
                show_message_dialog(self, 'Validation failed', 'Cannot start analysis:\n\n' + '\n'.join(validation_errors[:10]), kind='warning')
                self._analysis_in_progress = False
                self._reset_analyze_btn_idle()
                return
            self._log_analysis_run_params(selected_files=list(selected_files))
            self.current_task_index = 0
            self._start_analysis_progress(len(self.folder_fi_tasks), f'Analyzing {len(self.folder_fi_tasks)} task(s)...')
            self._process_next_folder_file()
            return
        input_file = Path(self.file_path.text()).resolve()
        if getattr(self, 'use_filtered_masks_cb', None) is not None and self.use_filtered_masks_cb.isChecked():
            try:
                from bactomate.gui.artifact_paths import resolve_filtered_mask_path
                name_lower = input_file.name.lower()
                if name_lower.endswith(('_masks.tif', '_masks.tiff')) and '_mask_filtered' not in name_lower:
                    resolved = resolve_filtered_mask_path(input_file)
                    if resolved is not None:
                        input_file = resolved
            except Exception:
                pass
        mask_path = input_file
        base_dir = base_dir_for_mask(mask_path)
        stem = stem_from_mask_path(mask_path)
        analysis_dir = analysis_dir_from_base(base_dir)
        out = analysis_dir / f'{stem}_analysis.xlsx'
        if self.analyze_fi.isChecked() and self._is_match_by_name_mode():
            fi_paths = self._resolve_fi_inputs_for_mask(mask_path)
            self.folder_files = [mask_path]
            self.folder_results = []
            self.folder_output_paths = []
            self.folder_fi_tasks = []
            if not fi_paths:
                self.folder_fi_tasks.append((mask_path, None, out))
            else:
                for fi_path in fi_paths:
                    suffix = self._fi_output_suffix(fi_path, stem)
                    self.folder_fi_tasks.append((mask_path, fi_path, analysis_dir / f'{stem}_analysis{suffix}.xlsx'))
            self._analysis_in_progress = True
            self._cancel_analysis_requested = False
            self._set_analyze_btn_cancel_mode(0)
            self.current_task_index = 0
            self._log_analysis_run_params(selected_files=[mask_path])
            self._start_analysis_progress(len(self.folder_fi_tasks), f'Analyzing {len(self.folder_fi_tasks)} task(s)...')
            self._process_next_folder_file()
            return
        fi_path = None
        if self.analyze_fi.isChecked() and (not self._is_match_by_name_mode()):
            fi_path = getattr(self, '_single_fi_path', None)
        self._analysis_in_progress = True
        self._cancel_analysis_requested = False
        try:
            self._set_analyze_btn_cancel_mode(0)
        except Exception:
            pass
        self._log_analysis_run_params(selected_files=[mask_path])
        self._run_one_analysis_task(mask_path, fi_path, out)

    def _build_analysis_run_manifest(self, *, mask_path: Path, fi_path: Path | None, fi_paths: list[Path] | None=None, out_path: Path, task_notes: list[str] | None=None, foci_outcome: Any | None=None, foci_available_not_imported: list[Any] | None=None) -> dict[str, Any]:
        from bactomate.gui.analysis_workbook import default_manifest_timestamp
        manifest: dict[str, Any] = {'Analysis date': default_manifest_timestamp(), 'Filtered mask file used': str(Path(mask_path).resolve()), 'Output workbook': str(Path(out_path).resolve()), 'Use filtered masks': 'Yes' if bool(getattr(self, 'use_filtered_masks_cb', None) and self.use_filtered_masks_cb.isChecked()) else 'No', 'Analyze FI': 'Yes' if bool(getattr(self, 'analyze_fi', None) and self.analyze_fi.isChecked()) else 'No', 'Include foci from detection CSV': 'Yes' if bool(getattr(self, 'include_foci_from_detection', None) and self.include_foci_from_detection.isChecked()) else 'No', 'Colocalization sheets': 'Yes' if bool(getattr(self, 'coloc_calc_all_pairs', None) and self.coloc_calc_all_pairs.isChecked()) else 'No', 'Per-cell detail sheet (--detailed)': 'Yes' if bool(getattr(self, 'detailed', None) and self.detailed.isChecked()) else 'No (may still be forced on for foci/coloc)', 'Analysis visualization': 'Yes (multi-panel TIFF)' if bool(getattr(self, 'create_vis', None) and self.create_vis.isChecked()) else 'No', 'Pixel size (µm)': str(float(self.pixel_size.value())) if hasattr(self, 'pixel_size') else ''}
        try:
            import bactomate
            manifest['BactoMate version'] = str(getattr(bactomate, '__version__', ''))
        except Exception:
            pass
        fi_inputs: list[Path] = []
        if fi_paths:
            fi_inputs = [Path(p) for p in fi_paths]
        elif fi_path is not None:
            fi_inputs = [Path(fi_path)]
        if fi_inputs:
            if len(fi_inputs) == 1:
                try:
                    manifest['FI image'] = str(fi_inputs[0].resolve())
                except Exception:
                    manifest['FI image'] = str(fi_inputs[0])
            else:
                stem = stem_from_mask_path(mask_path)
                lines: list[str] = []
                for p in fi_inputs:
                    tag = (self._fi_output_suffix(p, stem) or '').lstrip('_') or p.stem
                    try:
                        lines.append(f'{tag}: {p.resolve()}')
                    except Exception:
                        lines.append(f'{tag}: {p}')
                manifest['FI images'] = '\n'.join(lines)
        if getattr(self, 'analyze_fi', None) is not None and self.analyze_fi.isChecked():
            manifest['FI channel mode'] = 'Match by name' if self._is_match_by_name_mode() else 'One file for all masks'
            if getattr(self, 'analyze_fi_contour', None) and self.analyze_fi_contour.isChecked():
                manifest['Contour FI'] = f'Yes ({int(self.fi_contour_thickness.value())} px)'
            else:
                manifest['Contour FI'] = 'No'
        if getattr(self, 'include_foci_from_detection', None) and self.include_foci_from_detection.isChecked():
            if foci_outcome is not None:
                from bactomate.gui.analysis_workbook import outcome_to_manifest_entries
                manifest.update(outcome_to_manifest_entries(foci_outcome))
            else:
                manifest['Foci analysis result'] = 'Requested but CSV integration did not run'
        elif foci_available_not_imported:
            lines = [f'{a.channel_label}: {Path(a.csv_path).resolve()} ({a.csv_folder})' for a in foci_available_not_imported]
            manifest['Foci analysis result'] = 'Not imported — enable “Include foci from segmentation CSV” under Foci analysis and run Analyze again'
            manifest['Foci CSV available (not imported)'] = '\n'.join(lines)
        if getattr(self, 'coloc_calc_all_pairs', None) and self.coloc_calc_all_pairs.isChecked():
            manifest['Coloc max distance (px)'] = str(float(self.coloc_merge_max_distance.value()))
            manifest['Coloc threshold (px)'] = str(float(self.coloc_merge_threshold.value()))
            manifest['Coloc same cell only'] = 'Yes' if self.coloc_merge_same_cell_only.isChecked() else 'No'
        if task_notes:
            multi_output = len(task_notes) > 1 or any((' + FI ' in n for n in task_notes))
            if multi_output:
                manifest['Workbook outputs'] = '\n'.join(task_notes)
        return manifest

    def _run_post_analysis_workbook_steps(self, *, mask_path: Path, fi_path: Path | None, fi_paths: list[Path] | None=None, out_path: Path | None, task_notes: list[str] | None=None) -> None:
        """Foci CSV import, colocalization sheets, Run Info manifest (same workbook)."""
        if out_path is None:
            return
        out_path = Path(out_path)
        if not out_path.is_file():
            return
        mask_path = Path(mask_path)
        prefer_filtered = bool(getattr(self, 'use_filtered_masks_cb', None) and self.use_filtered_masks_cb.isChecked())
        foci_outcome = None
        foci_available_not_imported = None
        include_foci = bool(getattr(self, 'include_foci_from_detection', None) and self.include_foci_from_detection.isChecked())
        if not include_foci:
            try:
                from bactomate.gui.analysis_workbook import discover_foci_csv_artifacts_for_mask
                found = discover_foci_csv_artifacts_for_mask(self, mask_path, prefer_filtered=prefer_filtered)
                if found:
                    foci_available_not_imported = found
                    if hasattr(self.parent_window, 'log_output'):
                        self.parent_window.log_output('[Analyze] Foci CSV data was found but not imported. Enable “Include foci from segmentation CSV” under Foci analysis and run Analyze again to add Foci / per-cell columns.\n')
                    from bactomate.gui.widgets import show_confirm_dialog
                    if show_confirm_dialog(self, title='Import foci data?', message='Spotiflow foci CSV data was found for this mask, but “Include foci from segmentation CSV” is turned off.\n\nImport foci counts and intensities into this workbook now?', confirm_label='Yes, import', cancel_label='No', default_confirm=True):
                        include_foci = True
            except Exception:
                get_logger().debug('Foci artifact discovery failed', exc_info=True)
        if include_foci:
            self._set_post_analysis_progress('Importing foci into workbook…', percent=90)
            try:
                from bactomate.gui.analysis_workbook import integrate_foci_csv_into_workbook
                foci_outcome = integrate_foci_csv_into_workbook(self, mask_path, out_path, prefer_filtered=prefer_filtered)
                foci_available_not_imported = None
                if hasattr(self.parent_window, 'log_output'):
                    for line in foci_outcome.messages:
                        self.parent_window.log_output(line)
            except Exception as exc:
                get_logger().exception('Foci CSV integration failed')
                if hasattr(self.parent_window, 'log_output'):
                    self.parent_window.log_output(f'[Foci] Integration failed: {exc}\n')
        if getattr(self, 'coloc_calc_all_pairs', None) is not None and self.coloc_calc_all_pairs.isChecked():
            self._set_post_analysis_progress('Computing colocalization…', percent=93)
            try:
                from bactomate.gui.tabs._analyze_masks_colocalization import run_post_analysis_colocalization_for_masks
                run_post_analysis_colocalization_for_masks(self, [mask_path])
            except Exception:
                get_logger().debug('Post-analysis colocalization failed', exc_info=True)
        if bool(getattr(self, 'create_vis', None) and self.create_vis.isChecked()):
            self._set_post_analysis_progress('Writing analysis visualization…', percent=95)
            try:
                self._write_combined_analysis_visualization(mask_path=mask_path, analysis_path=out_path, prefer_filtered=prefer_filtered)
            except Exception as exc:
                get_logger().warning('Combined analysis visualization failed', exc_info=True)
                if hasattr(self.parent_window, 'log_output'):
                    self.parent_window.log_output(f'[Analyze] Analysis visualization failed: {exc}\n')
        self._set_post_analysis_progress('Writing Run Info…', percent=98)
        try:
            from bactomate.gui.analysis_workbook import write_run_info_sheet
            manifest = self._build_analysis_run_manifest(mask_path=mask_path, fi_path=fi_path, fi_paths=fi_paths, out_path=out_path, task_notes=task_notes, foci_outcome=foci_outcome, foci_available_not_imported=foci_available_not_imported)
            write_run_info_sheet(out_path, manifest)
            if hasattr(self.parent_window, 'log_output'):
                self.parent_window.log_output(f'[Analyze] Wrote Run Info sheet to {out_path.name}\n')
        except Exception:
            get_logger().debug('Run Info sheet write failed', exc_info=True)
        try:
            from bactomate.gui.analysis_workbook import write_master_per_cell_summary
            analysis_dir = analysis_dir_from_base(base_dir_for_mask(mask_path))

            def _master_log(msg: str) -> None:
                if hasattr(self.parent_window, 'log_output'):
                    self.parent_window.log_output(msg)
            write_master_per_cell_summary(analysis_dir, log_callback=_master_log)
        except Exception:
            get_logger().debug('Master per-cell summary write failed', exc_info=True)

    def _write_combined_analysis_visualization(self, *, mask_path: Path, analysis_path: Path, prefer_filtered: bool) -> None:
        """One horizontal multi-panel TIFF: Cells (phase) | each foci channel."""
        from bactomate.gui.analysis_visualization import write_combined_analysis_visualization
        mask_path = Path(mask_path)
        analysis_path = Path(analysis_path)
        base_dir = base_dir_for_mask(mask_path)
        analysis_dir = analysis_dir_from_base(base_dir)
        analysis_dir.mkdir(parents=True, exist_ok=True)
        out_vis = analysis_dir / f'{mask_path.stem}_analysis_visualization.tif'
        for legacy in (analysis_dir / f'{mask_path.stem}_filter_visualization.tif', *analysis_dir.glob(f'{mask_path.stem}_foci_visualization*.tif')):
            try:
                if legacy.is_file():
                    legacy.unlink()
            except Exception:
                pass
        ok = write_combined_analysis_visualization(self, mask_path=mask_path, output_path=out_vis, analysis_workbook=analysis_path if analysis_path.is_file() else None, prefer_filtered=prefer_filtered)
        if hasattr(self.parent_window, 'log_output'):
            if ok:
                self.parent_window.log_output(f'[Analyze] Wrote analysis visualization: {out_vis.name}\n')
            else:
                self.parent_window.log_output(f'[Analyze] Could not write analysis visualization ({out_vis.name}). Check the session log for details.\n')

    def _run_one_analysis_task(self, mask_path: Path, fi_path: Path | None, out_path: Path) -> None:
        (ok, _missing) = paths_exist([mask_path] + ([fi_path] if fi_path is not None else []))
        if not ok:
            self._analysis_in_progress = False
            try:
                self._reset_analyze_btn_idle()
            except Exception:
                pass
            show_message_dialog(self, 'Not Found', 'One or more inputs do not exist.', kind='warning')
            return
        if fi_path is not None:
            (ok, err) = mask_fi_spatial_shape_match(mask_path, fi_path)
            if not ok and err:
                self._analysis_in_progress = False
                try:
                    self._reset_analyze_btn_idle()
                except Exception:
                    pass
                show_message_dialog(self, 'Validation failed', err, kind='warning')
                return
        cmd_args = self._build_analyze_cell_areas_args(mask_path=mask_path, out_path=out_path, fi_path=fi_path)
        self.folder_output_paths = [str(out_path)]
        self._start_analysis_progress(1, f'Analyzing {mask_path.name}...')

        def _done(ok: bool, msg: str) -> None:
            try:
                self._analysis_worker = None
            except Exception:
                pass
            cancelled = bool(getattr(self, '_cancel_analysis_requested', False))
            try:
                if ok and out_path is not None and isinstance(out_path, Path) and (not out_path.exists()):
                    ok = False
                    msg = (msg or '').strip()
                    extra = f'Expected output was not created:\n{out_path}'
                    msg = (msg + '\n\n' + extra).strip() if msg else extra
            except Exception:
                pass
            if hasattr(self.parent_window, 'log_output') and msg:
                self.parent_window.log_output(msg + ('\n' if not msg.endswith('\n') else ''))
            if cancelled:
                self._analysis_in_progress = False
                try:
                    self._reset_analyze_btn_idle()
                except Exception:
                    pass
                self._close_analysis_progress(1)
                show_message_dialog(self, 'Analysis cancelled', 'Analysis was cancelled.', kind='info')
                return
            if ok:
                try:
                    self._set_post_analysis_progress('Finishing analysis…', percent=88)
                    self._run_post_analysis_workbook_steps(mask_path=Path(mask_path), fi_path=fi_path, out_path=Path(out_path) if out_path else None)
                except Exception:
                    get_logger().debug('Post-analysis workbook steps failed', exc_info=True)
            self._analysis_in_progress = False
            try:
                self._reset_analyze_btn_idle()
            except Exception:
                pass
            self._close_analysis_progress(1)
            self._show_finished_dialog_with_output_folder(title='Analysis complete' if ok else 'Analysis failed', message=msg or ('Done' if ok else 'Failed'), output_folder=out_path.parent if out_path is not None else None, icon='info' if ok else 'warning')
        t = WorkerThread('-m', cmd_args, parent=self)
        self._analysis_worker = t
        t.finished.connect(_done)
        if hasattr(t, 'progress'):
            t.progress.connect(self._handle_analysis_worker_progress)
        t.start()

    def _process_next_folder_file(self) -> None:
        if bool(getattr(self, '_cancel_analysis_requested', False)):
            self._analysis_in_progress = False
            self._close_analysis_progress()
            try:
                self._reset_analyze_btn_idle()
            except Exception:
                pass
            return
        tasks = getattr(self, 'folder_fi_tasks', []) or []
        idx = int(getattr(self, 'current_task_index', 0))
        if idx >= len(tasks):
            try:
                self._merge_fi_channels_into_one_excel()
            except Exception:
                pass
            try:
                results = getattr(self, 'folder_results', []) or []
                ok_masks: list[Path] = []
                by_mask: dict[Path, list[bool]] = {}
                tasks_by_mask: dict[Path, list[str]] = {}
                for ((m, fi_p, out_p), ok) in zip(tasks, results, strict=False):
                    mp = Path(m)
                    by_mask.setdefault(mp, []).append(bool(ok))
                    note = f'{mp.name}'
                    if fi_p is not None:
                        note += f' + FI {Path(fi_p).name}'
                    note += f' -> {Path(out_p).name}'
                    tasks_by_mask.setdefault(mp, []).append(note)
                for (m, oks) in by_mask.items():
                    if oks and all(oks):
                        ok_masks.append(Path(m))
                n_finish = len(ok_masks)
                for (fi, mp) in enumerate(ok_masks):
                    stem = stem_from_mask_path(mp)
                    out_p = analysis_dir_from_base(base_dir_for_mask(mp)) / f'{stem}_analysis.xlsx'
                    if not out_p.is_file():
                        continue
                    fi_paths_for_mask: list[Path] = []
                    for (m, fi_p, _o) in tasks:
                        if Path(m) == mp and fi_p is not None:
                            fi_paths_for_mask.append(Path(fi_p))
                    pct = min(99, 88 + int(10 * fi / max(1, n_finish)))
                    self._set_post_analysis_progress(f'Finishing {mp.name} ({fi + 1}/{n_finish})…', percent=pct)
                    self._run_post_analysis_workbook_steps(mask_path=mp, fi_path=fi_paths_for_mask[0] if len(fi_paths_for_mask) == 1 else None, fi_paths=fi_paths_for_mask if len(fi_paths_for_mask) > 1 else None, out_path=out_p, task_notes=tasks_by_mask.get(mp))
            except Exception:
                get_logger().debug('Folder post-analysis steps failed', exc_info=True)
            failures = []
            try:
                failures = getattr(self, '_folder_failures', []) or []
                if failures:
                    head = '\n\n'.join([f'- {name}\n{msg}'.strip() for (name, msg) in failures[:3]])
                    tail = ''
                    if len(failures) > 3:
                        tail = f'\n\n(+{len(failures) - 3} more; see log)'
                    self._analysis_finished_error_text = 'One or more analysis tasks failed or produced no output.\n\n' + head + tail
            except Exception:
                pass
            self._analysis_in_progress = False
            try:
                self._reset_analyze_btn_idle()
            except Exception:
                pass
            self._close_analysis_progress(len(tasks))
            output_folder = self._analysis_output_folder_from_paths(getattr(self, 'folder_output_paths', []) or [])
            if failures:
                self._show_finished_dialog_with_output_folder(title='Analysis finished with errors', message=getattr(self, '_analysis_finished_error_text', 'One or more analysis tasks failed or produced no output.'), output_folder=output_folder, icon='warning')
            else:
                self._show_finished_dialog_with_output_folder(title='Analysis complete', message=f'Analysis finished. {len(tasks)} task(s) completed.', output_folder=output_folder)
            return
        (mask_path, fi_path, out_path) = tasks[idx]
        self._update_analysis_progress(idx, f'Analyzing {Path(mask_path).name} ({idx + 1}/{len(tasks)})...')
        try:
            self._set_analyze_btn_cancel_mode(0)
        except Exception:
            pass
        cmd_args = self._build_analyze_cell_areas_args(mask_path=mask_path, out_path=out_path, fi_path=fi_path)
        th = WorkerThread('-m', cmd_args, parent=self)
        self._analysis_worker = th

        def _on_done(success: bool, message: str):
            try:
                self._analysis_worker = None
            except Exception:
                pass
            try:
                if success and out_path is not None and (not Path(out_path).exists()):
                    success = False
                    extra = f'Expected output was not created: {out_path}'
                    message = (str(message or '').strip() + '\n' + extra).strip()
            except Exception:
                pass
            getattr(self, 'folder_results', []).append(bool(success))
            getattr(self, 'folder_output_paths', []).append(str(out_path))
            if hasattr(self.parent_window, 'log_output') and message:
                self.parent_window.log_output(message + ('\n' if not message.endswith('\n') else ''))
            if not success:
                try:
                    lst = getattr(self, '_folder_failures', None)
                    if isinstance(lst, list):
                        lst.append((Path(mask_path).name, str(message or 'Failed').strip()))
                except Exception:
                    pass
            self.current_task_index = idx + 1
            self._update_analysis_progress(idx + 1, f'Finished {Path(mask_path).name} ({idx + 1}/{len(tasks)})...')
            self._process_next_folder_file()
        th.finished.connect(_on_done)
        if hasattr(th, 'progress'):
            th.progress.connect(self._handle_analysis_worker_progress)
        th.start()

    def _build_analyze_cell_areas_args(self, *, mask_path: Path, out_path: Path, fi_path: Path | None) -> list[str]:
        """
        Build CLI args for `bactomate.scripts.analyze_cell_areas` matching its argparse contract.

        Script signature:
          analyze_cell_areas masks --output OUT [--pixel-size ...] [--fi-image ...] ...
        """
        args: list[str] = [self._analyze_script_module, str(mask_path), '--output', str(out_path)]

        def _f(w, default: float) -> float:
            try:
                return float(w.value())
            except Exception:
                return float(default)

        def _opt_um2(w) -> float | None:
            try:
                v = float(w.value())
                return v if v > 0 else None
            except Exception:
                return None
        args += ['--pixel-size', str(_f(getattr(self, 'pixel_size', None), 0.108))]
        detailed_on = bool(getattr(self, 'detailed', None) is not None and self.detailed.isChecked())
        if getattr(self, 'include_foci_from_detection', None) is not None and self.include_foci_from_detection.isChecked():
            detailed_on = True
        if getattr(self, 'coloc_calc_all_pairs', None) is not None and self.coloc_calc_all_pairs.isChecked():
            detailed_on = True
        if detailed_on:
            args += ['--detailed']
        if getattr(self, 'analyze_fi', None) is not None and self.analyze_fi.isChecked() and (fi_path is not None):
            args += ['--fi-image', str(fi_path)]
            if getattr(self, 'analyze_fi_contour', None) is not None and self.analyze_fi_contour.isChecked():
                args += ['--fi-contour-thickness', str(int(self.fi_contour_thickness.value()))]
            else:
                args += ['--no-fi-contour']
        return args

    def _fi_output_suffix(self, fi_path: Path | None, stem: str) -> str:
        if fi_path is None:
            return ''
        try:
            tag = self._extract_channel_tag_for_mask_stem(fi_path.stem, stem) or ''
        except Exception:
            tag = ''
        if tag:
            return '_' + str(tag)
        s = fi_path.stem
        if s.startswith(stem + '_'):
            rem = s[len(stem) + 1:]
            if rem:
                return '_' + rem
        return ''

    def _match_log_from_folder_tasks(self) -> tuple[int, int, list[str]]:
        tasks = getattr(self, 'folder_fi_tasks', [])
        results = getattr(self, 'folder_results', [])
        if len(results) != len(tasks):
            return (0, 0, [])
        seen_stem: set[str] = set()
        lines: list[str] = []
        mask_stems: set[str] = set()
        n_tif = 0
        for (i, (mask_file, fi_path, _out)) in enumerate(tasks):
            if not results[i]:
                continue
            stem = mask_file.stem.replace('_masks', '')
            tif_name = fi_path.name if fi_path is not None else '—'
            if fi_path is not None:
                n_tif += 1
            mask_stems.add(stem)
            if stem not in seen_stem:
                seen_stem.add(stem)
                lines.append(f'{stem}_masks -> {tif_name}')
            else:
                lines.append(f'{stem}_masks with {tif_name}')
        return (len(mask_stems), n_tif, lines)

    def _build_combined_summary_excel(self, excel_paths, output_path, match_log=None, log_callback=None):

        def _log(msg: str) -> None:
            if log_callback:
                log_callback(msg)
        import re
        import pandas as pd
        from openpyxl import load_workbook
        from openpyxl.utils.dataframe import dataframe_to_rows

        def _sanitize_sheet_name(name: str, max_len: int=31) -> str:
            s = re.sub('[\\\\/:*?\\[\\]]', '_', str(name))
            return s[:max_len] if len(s) > max_len else s
        if not excel_paths:
            _log('Combined summary: no paths provided.\n')
            return False
        frames_long = []
        missing = []
        no_summary = []
        per_cell_cache: dict[Path, pd.DataFrame | None] = {}
        for p in excel_paths:
            p = Path(p)
            if not p.exists():
                missing.append(p.name)
                continue
            try:
                with pd.ExcelFile(p) as xl:
                    df = pd.read_excel(xl, sheet_name='Summary')
                    per_cell = None
                    try:
                        per_cell = pd.read_excel(xl, sheet_name='Per-Cell Data')
                    except Exception:
                        pass
                    per_cell_cache[p] = per_cell
            except Exception as e:
                no_summary.append((p.name, str(e)))
                continue
            if df is None or df.empty:
                no_summary.append((p.name, 'empty Summary sheet'))
                continue
            df = df.copy()
            df['File'] = p.stem.replace('_analysis', '').replace('_masks_analysis', '')
            frames_long.append(df)
        if missing:
            _log(f"Combined summary: {len(missing)} path(s) not found: {', '.join(missing[:5])}{(' ...' if len(missing) > 5 else '')}\n")
        if no_summary and (not frames_long):
            _log(f'Combined summary: no Summary sheet could be read. First failure: {no_summary[0][0]} — {no_summary[0][1]}\n')
        if not frames_long:
            return False
        long_df = pd.concat(frames_long, ignore_index=True)
        cols = [c for c in ['File', 'Frame'] if c in long_df.columns]
        cols += [c for c in long_df.columns if c not in cols]
        long_df = long_df[cols]
        wide_df = None
        wide_err = None
        try:
            if 'Frame' not in long_df.columns:
                wide_err = "Summary sheets have no 'Frame' column"
            else:
                frames = sorted(long_df['Frame'].dropna().unique().astype(int).tolist())
                if not frames:
                    wide_err = 'Summary sheets have no Frame values'
                else:
                    wide_data = {'Frame': frames}
                    for f in long_df['File'].unique().tolist():
                        sub = long_df[long_df['File'] == f].drop_duplicates(subset=['Frame'], keep='first').set_index('Frame')
                        for col in ['Total Cell Area (µm²)', 'Total Cell Area (Relative)', 'Number of Cells']:
                            if col in sub.columns:
                                wide_data[f'{f}_{col}'] = sub.reindex(frames)[col].values
                    wide_df = pd.DataFrame(wide_data)
        except Exception as e:
            wide_err = str(e)
            wide_df = None
        if wide_df is None or wide_df.empty:
            if wide_err:
                _log(f'Combined summary: wide table build failed — {wide_err}\n')
            return False
        out_path = Path(output_path)
        existing_log_lines: list[str] = []
        if out_path.exists() and match_log is not None:
            try:
                wb_existing = load_workbook(output_path, data_only=True)
                sheet_name_wide = 'Combined Summary (wide)'
                if sheet_name_wide in wb_existing.sheetnames:
                    ws_existing = wb_existing[sheet_name_wide]
                    for r in range(2, min(ws_existing.max_row + 1, 5000)):
                        val = ws_existing.cell(row=r, column=7).value
                        if val is None or not str(val).strip():
                            continue
                        s = str(val).strip()
                        if ' -> ' in s or ' with ' in s:
                            existing_log_lines.append(s)
                wb_existing.close()
            except Exception:
                pass
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = 'Combined Summary (wide)'
            for row in dataframe_to_rows(wide_df, index=False, header=True):
                ws.append(row)
            if match_log is not None:
                (n_masks, n_tif, lines) = match_log
                ws['G1'] = 'Logs'
                ws['G2'] = f'{n_masks} mask(s) and {n_tif} tif file(s) matched'
                rowi = 3
                for line in existing_log_lines:
                    ws.cell(row=rowi, column=7, value=line)
                    rowi += 1
                for line in lines:
                    ws.cell(row=rowi, column=7, value=line)
                    rowi += 1
            seen_sheet_names = set(wb.sheetnames)
            for p in excel_paths:
                p = Path(p)
                if not p.exists():
                    continue
                per_cell = per_cell_cache.get(p)
                if per_cell is None or per_cell.empty:
                    continue
                base_name = p.stem.replace('_analysis', '').replace('_masks_analysis', '')
                sheet_name = _sanitize_sheet_name(base_name)
                if sheet_name in seen_sheet_names:
                    stem2 = sheet_name[:28] if len(sheet_name) > 28 else sheet_name
                    n = 1
                    while f'{stem2}_{n}' in seen_sheet_names:
                        n += 1
                    sheet_name = f'{stem2}_{n}'
                seen_sheet_names.add(sheet_name)
                new_ws = wb.create_sheet(title=sheet_name)
                for row in dataframe_to_rows(per_cell, index=False, header=True):
                    new_ws.append(row)
            wb.save(output_path)
        except Exception as e:
            _log(f'Combined summary: write failed — {e}\n')
            return False
        return True

    def _merge_fi_channels_into_one_excel(self) -> None:
        import pandas as pd
        tasks = getattr(self, 'folder_fi_tasks', [])
        results = getattr(self, 'folder_results', [])
        if len(results) != len(tasks):
            return

        def _col_key(raw: str) -> str:
            txt = str(raw or '').strip()
            if not txt:
                return 'FI'
            out = []
            for ch in txt:
                out.append(ch if ch.isalnum() or ch == '_' else '_')
            key = ''.join(out).strip('_')
            while '__' in key:
                key = key.replace('__', '_')
            return key or 'FI'
        by_stem: dict[str, list[tuple[Path, Path | None, str]]] = {}
        for (i, (mask_file, fi_path, output_path)) in enumerate(tasks):
            if not results[i] or not output_path.exists():
                continue
            stem = mask_file.stem.replace('_masks', '')
            by_stem.setdefault(stem, [])
            channel_label = 'FI'
            if fi_path is not None:
                suffix = self._fi_output_suffix(fi_path, stem)
                channel_label = (suffix[1:] if suffix.startswith('_') else 'FI') or 'FI'
            by_stem[stem].append((output_path, fi_path, channel_label))
        for (stem, group) in by_stem.items():
            phase_entries = [(p, fp, lbl) for (p, fp, lbl) in group if fp is None]
            fi_entries = [(p, fp, lbl) for (p, fp, lbl) in group if fp is not None]
            if not fi_entries:
                continue
            analysis_dir = group[0][0].parent
            merged_path = analysis_dir / f'{stem}_analysis.xlsx'
            pre_coloc_sheets: list[str] = []
            try:
                if merged_path.exists():
                    from openpyxl import load_workbook
                    wb_pre = load_workbook(merged_path, read_only=True)
                    pre_coloc_sheets = [s for s in list(wb_pre.sheetnames) if isinstance(s, str) and s.startswith('Colocalization ')]
                    wb_pre.close()
            except Exception:
                pre_coloc_sheets = []
            try:
                if pre_coloc_sheets and hasattr(self.parent_window, 'log_output'):
                    self.parent_window.log_output(f'[Analyze/Merge] {merged_path.name} currently has {len(pre_coloc_sheets)} Colocalization sheet(s); FI merge rewrites workbook before post-colocalization append.\n')
            except Exception:
                pass
            try:
                summary_df = None
                merged = None
                first_excel = phase_entries[0][0] if phase_entries else fi_entries[0][0]
                for (_p, _fp, _lbl) in fi_entries:
                    try:
                        summary_df = pd.read_excel(_p, sheet_name='Summary')
                        break
                    except Exception:
                        continue
                if summary_df is None and phase_entries:
                    try:
                        summary_df = pd.read_excel(phase_entries[0][0], sheet_name='Summary')
                    except Exception:
                        pass
                try:
                    head = pd.read_excel(first_excel, sheet_name='Per-Cell Data', nrows=0)
                    key_cols = ['Frame', 'Cell_ID'] if 'Cell_ID' in head.columns else ['Frame', 'Cell Label']
                except Exception:
                    key_cols = ['Frame', 'Cell Label']
                fi_col_names = ['Mean FI Intensity', 'Mean FI Intensity (whole cell)', 'Contour FI Intensity', 'FI Ratio (Contour/Mean)']
                fi_col_bases = list(fi_col_names)

                def _strip_bracket_tag(col: str) -> tuple[str, str | None]:
                    """
                    Parse columns of the form "Base [TAG]" -> ("Base", "TAG").
                    Returns (original, None) when no bracket tag is present.
                    """
                    s = str(col or '')
                    if '[' not in s or ']' not in s:
                        return (s, None)
                    i = s.rfind('[')
                    j = s.rfind(']')
                    if i < 0 or j < 0 or j < i:
                        return (s, None)
                    base = s[:i].rstrip()
                    tag = s[i + 1:j].strip()
                    if not base or not tag:
                        return (s, None)
                    if j != len(s) - 1:
                        return (s, None)
                    return (base, tag)

                def _canonicalize_fi_columns(df: 'pd.DataFrame') -> tuple['pd.DataFrame', dict[str, str]]:
                    """
                    Return (df2, colmap) where df2 has FI columns renamed from e.g.
                    "Contour FI Intensity [GFP]" -> "Contour FI Intensity" (base name only),
                    so downstream merging can consistently suffix with (channel_label).
                    """
                    if df is None or df.empty:
                        return (df, {})
                    colmap: dict[str, str] = {}
                    for c in list(df.columns):
                        if not isinstance(c, str):
                            continue
                        (base, _tag) = _strip_bracket_tag(c)
                        if base in fi_col_bases and base != c:
                            colmap[c] = base
                    if not colmap:
                        return (df, {})
                    return (df.rename(columns=colmap), colmap)
                if phase_entries:
                    phase_path = phase_entries[0][0]
                    try:
                        df_phase = pd.read_excel(phase_path, sheet_name='Per-Cell Data')
                        if df_phase is not None and (not df_phase.empty):
                            merged = df_phase.copy()
                    except Exception:
                        pass
                for (excel_path, _fp, channel_label) in fi_entries:
                    try:
                        df = pd.read_excel(excel_path, sheet_name='Per-Cell Data')
                    except Exception:
                        df = None
                    if df is None or df.empty:
                        continue
                    (df, _fi_colmap) = _canonicalize_fi_columns(df)
                    has_fi = [c for c in fi_col_names if c in df.columns]
                    rename = {c: f'{c} ({channel_label})' for c in has_fi}
                    sub = df.rename(columns=rename)
                    merge_on = [c for c in key_cols if merged is not None and c in merged.columns and (c in sub.columns)]

                    def is_fi_or_foci_col(c: str) -> bool:
                        if c in key_cols:
                            return False
                        up = str(c).upper()
                        if 'FI' in str(c) or 'FOCI' in up or c == 'Number of Foci':
                            return True
                        if str(c).startswith('foci_ids_') or str(c).endswith('_foci_count'):
                            return True
                        return False
                    if merged is None:
                        merged = sub.copy()
                    else:
                        fi_new = [c for c in sub.columns if c not in merge_on and is_fi_or_foci_col(c) and (c not in merged.columns)]
                        if merge_on and fi_new:
                            merged = merged.merge(sub[merge_on + fi_new], on=merge_on, how='outer')
                if summary_df is None:
                    continue
                try:
                    if merged is not None and (not merged.empty) and ('Frame' in merged.columns):
                        summary_df = summary_df.copy()
                        used_cols = set((str(c) for c in summary_df.columns))
                        frame_series = pd.to_numeric(merged['Frame'], errors='coerce')
                        merged_with_frame = merged.copy()
                        merged_with_frame['_frame_int'] = frame_series.astype('Int64')

                        def _add_frame_mean(col_name: str, out_col: str) -> None:
                            nonlocal summary_df
                            if out_col in used_cols:
                                return
                            if col_name not in merged_with_frame.columns:
                                return
                            s = pd.to_numeric(merged_with_frame[col_name], errors='coerce')
                            tmp = pd.DataFrame({'Frame': merged_with_frame['_frame_int'], '__v': s})
                            tmp = tmp.dropna(subset=['Frame'])
                            if tmp.empty:
                                return
                            means = tmp.groupby('Frame', dropna=True)['__v'].mean()
                            try:
                                sf = pd.to_numeric(summary_df['Frame'], errors='coerce').astype('Int64')
                            except Exception:
                                return
                            summary_df[out_col] = sf.map(lambda f: float(means.get(int(f), np.nan)) if f is not None and pd.notna(f) else np.nan)
                            used_cols.add(out_col)
                        import numpy as np
                        for (_excel_path, _fp, ch) in fi_entries:
                            for base in fi_col_bases:
                                src = f'{base} ({ch})'
                                _add_frame_mean(src, src)
                except Exception:
                    pass
                if merged is not None and (not merged.empty):
                    simple_foci_cols: set[str] = set()
                    for col in list(merged.columns):
                        if not isinstance(col, str):
                            continue
                        if col.startswith('foci_ids_'):
                            ch = col[len('foci_ids_'):].strip() or 'FI'
                            alias = f'foci id [{ch}]'
                            merged[alias] = merged[col]
                            simple_foci_cols.add(alias)
                        elif col.endswith('_foci_count'):
                            ch = col[:-len('_foci_count')].strip() or 'FI'
                            alias = f'num foci [{ch}]'
                            merged[alias] = merged[col]
                            simple_foci_cols.add(alias)

                    def _is_foci_clutter_col(c: str) -> bool:
                        if c in simple_foci_cols:
                            return False
                        if c.startswith('num_coloc_foci_'):
                            return False
                        up = c.upper()
                        if c.startswith('foci_ids_') or c.endswith('_foci_count'):
                            return True
                        if 'FOCI' in up:
                            return True
                        return False
                    drop_cols = [c for c in list(merged.columns) if isinstance(c, str) and _is_foci_clutter_col(c)]
                    if drop_cols:
                        merged = merged.drop(columns=drop_cols, errors='ignore')
                foci_by_channel: dict[str, list[pd.DataFrame]] = {}
                for (excel_path, _fp, ch_lbl) in fi_entries:
                    try:
                        df_foci = pd.read_excel(excel_path, sheet_name='Foci')
                    except Exception:
                        df_foci = None
                    if df_foci is None or df_foci.empty:
                        continue
                    df_foci = df_foci.copy()
                    if 'channel_name' not in df_foci.columns:
                        df_foci['channel_name'] = ch_lbl
                    foci_by_channel.setdefault(ch_lbl, []).append(df_foci)

                def _safe_excel_sheet_name(base: str, used: set[str]) -> str:
                    bad = '\\/:*?[]'
                    safe = ''.join(('_' if ch in bad else ch for ch in str(base))).strip() or 'Sheet'
                    safe = safe[:31]
                    if safe not in used:
                        used.add(safe)
                        return safe
                    stem_name = safe[:28] if len(safe) > 28 else safe
                    n = 2
                    cand = f'{stem_name}_{n}'
                    while cand in used:
                        n += 1
                        cand = f'{stem_name}_{n}'
                    used.add(cand)
                    return cand
                used_sheet_names = {'Summary', 'Per-Cell Data'}
                with pd.ExcelWriter(merged_path, engine='openpyxl') as writer:
                    summary_df.to_excel(writer, sheet_name='Summary', index=False)
                    if merged is not None and (not merged.empty):
                        merged.to_excel(writer, sheet_name='Per-Cell Data', index=False)
                    for ch in sorted(foci_by_channel.keys(), key=lambda s: str(s).lower()):
                        parts = foci_by_channel.get(ch) or []
                        if not parts:
                            continue
                        ch_df = pd.concat(parts, ignore_index=True, sort=False)
                        sheet_name = _safe_excel_sheet_name(f'Foci {ch}', used_sheet_names)
                        ch_df.to_excel(writer, sheet_name=sheet_name, index=False)
                try:
                    from openpyxl import load_workbook
                    wb_post = load_workbook(merged_path, read_only=True)
                    post_names = list(wb_post.sheetnames)
                    post_coloc = [s for s in post_names if isinstance(s, str) and s.startswith('Colocalization ')]
                    wb_post.close()
                    if hasattr(self.parent_window, 'log_output'):
                        self.parent_window.log_output(f'[Analyze/Merge] Wrote {merged_path.name}: sheets={len(post_names)}, coloc_sheets_after_merge={len(post_coloc)}.\n')
                        if pre_coloc_sheets and (not post_coloc):
                            self.parent_window.log_output('[Analyze/Merge] Existing Colocalization sheets were removed by merge rewrite; they should be re-added by post-analysis colocalization step.\n')
                except Exception:
                    pass
            except Exception:
                get_logger().exception('Merge FI channels failed')
                continue

    def get_session_params(self) -> dict[str, Any]:
        d = super().get_session_params()
        d.update({'analyze_fi': self._safe_checked(getattr(self, 'analyze_fi', None), default=False), 'fi_match_by_name': self._safe_checked(getattr(self, 'fi_mode_match_by_name', None), default=True), 'single_fi_path': str(self._single_fi_path) if getattr(self, '_single_fi_path', None) is not None else '', 'analyze_fi_contour': self._safe_checked(getattr(self, 'analyze_fi_contour', None), default=False), 'fi_contour_thickness': int(self._safe_value(getattr(self, 'fi_contour_thickness', None), default=1)), 'use_filtered_masks': self._safe_checked(getattr(self, 'use_filtered_masks_cb', None), default=True), 'coloc_merge_max_distance': self._safe_value(getattr(self, 'coloc_merge_max_distance', None), default=8.0), 'coloc_merge_threshold': self._safe_value(getattr(self, 'coloc_merge_threshold', None), default=4.0), 'coloc_merge_same_cell_only': self._safe_checked(getattr(self, 'coloc_merge_same_cell_only', None), default=True), 'coloc_calc_all_pairs': self._safe_checked(getattr(self, 'coloc_calc_all_pairs', None), default=False), 'include_foci_from_detection': self._safe_checked(getattr(self, 'include_foci_from_detection', None), default=False)})
        return d

    def set_session_params(self, data: dict[str, Any]) -> None:
        super().set_session_params(data)
        apply_to_widgets(data, [('analyze_fi', self.analyze_fi), ('analyze_fi_contour', self.analyze_fi_contour), ('fi_contour_thickness', self.fi_contour_thickness), ('use_filtered_masks', self.use_filtered_masks_cb), ('coloc_merge_max_distance', self.coloc_merge_max_distance), ('coloc_merge_threshold', self.coloc_merge_threshold), ('coloc_merge_same_cell_only', self.coloc_merge_same_cell_only), ('coloc_calc_all_pairs', self.coloc_calc_all_pairs), ('include_foci_from_detection', self.include_foci_from_detection)])
        fi_match = data.get('fi_match_by_name', data.get('batch_fi', True))
        self._safe_set_checked(getattr(self, 'fi_mode_match_by_name', None), bool(fi_match))
        self._safe_set_checked(getattr(self, 'fi_mode_one_file', None), not bool(fi_match))
        if 'single_fi_path' in data and data['single_fi_path']:
            try:
                self._single_fi_path = Path(data['single_fi_path'])
            except Exception:
                self._single_fi_path = None
        self._update_fi_ui_state()
