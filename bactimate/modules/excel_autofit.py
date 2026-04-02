"""
Heuristic column widths for workbooks written with pandas + openpyxl (Excel’s default is uniform narrow columns).
"""
from __future__ import annotations


def _display_len(value: object) -> int:
    if value is None:
        return 0
    s = str(value)
    if "\n" in s:
        s = s.split("\n", 1)[0]
    return len(s)


def autofit_openpyxl_workbook(
    workbook,
    *,
    min_width: float = 8.0,
    max_width: float = 60.0,
    max_rows: int = 1000,
) -> None:
    """
    Adjust each column’s width from the longest sampled cell in that column (header + data).
    Capped so huge text fields do not make columns absurdly wide.
    """
    try:
        from openpyxl.utils import get_column_letter
    except ImportError:
        return

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        if ws.max_column == 0:
            continue
        last_row = min(ws.max_row, max(1, max_rows))
        for col_idx in range(1, ws.max_column + 1):
            letter = get_column_letter(col_idx)
            longest = min_width
            for row_idx in range(1, last_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                L = float(_display_len(cell.value)) + 1.8
                if L > longest:
                    longest = L
            ws.column_dimensions[letter].width = max(min_width, min(max_width, longest))
