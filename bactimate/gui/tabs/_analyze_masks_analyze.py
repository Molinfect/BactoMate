from __future__ import annotations

import subprocess
import sys
from pathlib import Path
from typing import Any, Optional

import numpy as np
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QButtonGroup,
    QCheckBox,
    QDoubleSpinBox,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QRadioButton,
    QSpinBox,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from bactimate.gui.analysis_paths import analysis_dir_from_base, base_dir_for_mask, stem_from_mask_path
from bactimate.gui.constants import PROJECT_ROOT
from bactimate.gui.preferences import get_default_browse_dir
from bactimate.gui.session_params import apply_to_widgets
from bactimate.gui.widgets import DragDropGroupBox, create_info_icon, show_file_selection_dialog
from bactimate.gui.workers import WorkerThread
from bactimate.gui.log_config import get_logger
from bactimate.modules.validation import mask_fi_spatial_shape_match, paths_exist

from bactimate.gui.tabs._analyze_masks_base import AnalyzeMasksBase


class AnalyzeTab(AnalyzeMasksBase):
    """Section 3: Analyze (FI / contour / colocalization reporting; writes *_analysis*.xlsx)."""

    # ---- base hooks ---------------------------------------------------
    def _stage_key(self) -> str:
        return "analyze"

    def _file_group_title(self) -> str:
        return "Analyze artifacts"

    def _preview_button_label(self) -> str:
        # Analyze section does not have a filtering preview
        return "Preview"

    def _primary_button_label(self) -> str:
        return "Analyze File/Folder"

    def _primary_button_tooltip(self) -> str:
        return "Run analysis and write *_analysis.xlsx outputs"

    def _stage_note_text(self) -> str:
        return "Analysis runs from this section. If you need to tune filtering settings, use the Filtering section above."

    def _wire_stage_buttons(self) -> None:
        # No preview button in Analyze section
        try:
            self.preview_btn.setVisible(False)
        except Exception:
            pass
        self.analyze_btn.clicked.connect(self.analyze_masks)

    # ---- file group override (use filtered masks + live hint) --------
    def _build_file_group(self) -> QWidget:
        file_group = DragDropGroupBox(
            self._file_group_title(),
            parent=self,
            on_drop_callback=self.handle_dropped_path_analyze,
        )
        file_layout = QHBoxLayout()
        file_layout.addWidget(QLabel("Path:"))
        file_layout.addWidget(
            create_info_icon(
                "Path to a mask TIFF file or folder of mask TIFFs to analyze.\n\n"
                "When enabled, analysis prefers `*_mask_filtered.tif` artifacts and falls back to `*_masks.tif`.",
                self,
            )
        )
        self.file_path = QLineEdit()
        self.file_path.setPlaceholderText("Select a mask TIFF file or folder... (or drag and drop)")
        self.file_path.setAcceptDrops(False)
        file_btn = QPushButton("Browse File...")
        file_btn.setFixedWidth(120)
        file_btn.clicked.connect(self.select_file)
        folder_btn = QPushButton("Browse Folder...")
        folder_btn.setFixedWidth(130)
        folder_btn.clicked.connect(self.select_folder)
        file_layout.addWidget(self.file_path, 1)
        file_layout.addWidget(file_btn, 0)
        file_layout.addWidget(folder_btn, 0)

        self.use_filtered_masks_cb = QCheckBox("Use filtered masks if available")
        self.use_filtered_masks_cb.setChecked(True)
        self.use_filtered_masks_cb.setToolTip(
            "When checked, analysis will prefer `*_mask_filtered.tif` artifacts.\n"
            "If none are found, it falls back to `*_masks.tif`."
        )
        self.filtered_masks_hint = QLabel("")
        self.filtered_masks_hint.setStyleSheet("color: #9e9e9e; font-size: 9pt;")

        row2 = QHBoxLayout()
        row2.addWidget(self.use_filtered_masks_cb)
        row2.addSpacing(8)
        row2.addWidget(self.filtered_masks_hint, 1)
        row2.addStretch()

        outer = QVBoxLayout()
        outer.setContentsMargins(0, 4, 0, 0)
        outer.setSpacing(4)
        outer.addLayout(file_layout)
        outer.addLayout(row2)
        file_group.setLayout(outer)

        def _refresh_hint() -> None:
            try:
                txt = (self.file_path.text() or "").strip()
                if not txt:
                    self.filtered_masks_hint.setText("")
                    return
                p = Path(txt)
                if not p.exists() or not self.use_filtered_masks_cb.isChecked():
                    self.filtered_masks_hint.setText("")
                    return
                folder = p if p.is_dir() else p.parent
                n_filt = len(list(folder.glob("*_mask_filtered.tif"))) + len(list(folder.glob("*_mask_filtered.tiff")))
                if n_filt:
                    self.filtered_masks_hint.setText(f"({n_filt} filtered mask(s) found)")
                else:
                    self.filtered_masks_hint.setText("(no filtered masks found; will fall back)")
            except Exception:
                self.filtered_masks_hint.setText("")

        self.file_path.textChanged.connect(lambda: _refresh_hint())
        self.use_filtered_masks_cb.toggled.connect(lambda _v: _refresh_hint())
        _refresh_hint()
        return file_group

    # ---- analyze-stage panel (FI + contour + colocalization controls) -
    def _build_stage_analysis_subpanel(self) -> Optional[QWidget]:
        outer = QWidget()
        outer_layout = QVBoxLayout(outer)
        outer_layout.setContentsMargins(0, 6, 0, 0)
        outer_layout.setSpacing(6)

        # Analyze FI toggle
        fi_toggle_row = QHBoxLayout()
        self.analyze_fi = QCheckBox("Analyze FI")
        self.analyze_fi.setChecked(False)
        fi_toggle_row.addWidget(self.analyze_fi)
        fi_toggle_row.addWidget(
            create_info_icon(
                "Analyze fluorescence intensity (FI) from a separate channel.\n\n"
                "Match by name will search for `stem_fi.tif` / `stem_fi_*` next to the mask.",
                parent_widget=self,
            )
        )
        fi_toggle_row.addStretch()
        outer_layout.addLayout(fi_toggle_row)

        fi_group = QGroupBox("FI channels")
        fi_layout = QVBoxLayout()
        self.fi_mode_one_file = QRadioButton("One FI file for all masks")
        self.fi_mode_match_by_name = QRadioButton("Match FI channel(s) by name")
        self.fi_mode_match_by_name.setChecked(True)
        self.fi_mode_group = QButtonGroup(self)
        self.fi_mode_group.addButton(self.fi_mode_one_file)
        self.fi_mode_group.addButton(self.fi_mode_match_by_name)
        mode_row = QHBoxLayout()
        mode_row.addWidget(self.fi_mode_one_file)
        mode_row.addWidget(self.fi_mode_match_by_name)
        mode_row.addStretch()
        fi_layout.addLayout(mode_row)

        self._single_fi_path: Optional[Path] = None
        self.fi_select_btn = QPushButton("Select FI file…")
        self.fi_select_btn.setFixedWidth(140)
        self.fi_select_btn.clicked.connect(self._select_fi_file)
        fi_layout.addWidget(self.fi_select_btn)

        self.fi_feedback = QTextEdit()
        self.fi_feedback.setReadOnly(True)
        self.fi_feedback.setMaximumHeight(140)
        self.fi_feedback.setPlaceholderText("Set mask path and choose FI mode to see what will be analyzed.")
        fi_layout.addWidget(self.fi_feedback)

        contour_group = QGroupBox("Contour FI")
        contour_row = QHBoxLayout()
        self.analyze_fi_contour = QCheckBox("Contour FI")
        self.analyze_fi_contour.setChecked(False)
        self.fi_contour_thickness = QSpinBox()
        self.fi_contour_thickness.setRange(1, 10)
        self.fi_contour_thickness.setValue(1)
        self.fi_contour_thickness.setEnabled(False)
        self.fi_contour_thickness.setFixedWidth(80)
        contour_row.addWidget(self.analyze_fi_contour)
        contour_row.addWidget(QLabel("Contour thickness (pixels):"))
        contour_row.addWidget(self.fi_contour_thickness)
        contour_row.addStretch()
        contour_group.setLayout(contour_row)

        fi_layout.addWidget(contour_group)
        fi_group.setLayout(fi_layout)
        outer_layout.addWidget(fi_group)

        coloc_group = QGroupBox("Colocalization (multi-FI merge)")
        coloc_outer = QVBoxLayout()
        row = QHBoxLayout()
        row.addWidget(QLabel("Max distance (px):"))
        self.coloc_merge_max_distance = QDoubleSpinBox()
        self.coloc_merge_max_distance.setRange(0.1, 500.0)
        self.coloc_merge_max_distance.setDecimals(1)
        self.coloc_merge_max_distance.setValue(8.0)
        row.addWidget(self.coloc_merge_max_distance)
        row.addWidget(QLabel("Threshold (px):"))
        self.coloc_merge_threshold = QDoubleSpinBox()
        self.coloc_merge_threshold.setRange(0.1, 500.0)
        self.coloc_merge_threshold.setDecimals(1)
        self.coloc_merge_threshold.setValue(4.0)
        row.addWidget(self.coloc_merge_threshold)
        self.coloc_merge_same_cell_only = QCheckBox("Same cell only")
        self.coloc_merge_same_cell_only.setChecked(True)
        row.addWidget(self.coloc_merge_same_cell_only)
        row.addStretch()
        coloc_outer.addLayout(row)
        self.coloc_calc_all_pairs = QCheckBox(
            "After successful analysis, add one colocalization sheet per *_foci channel pair to *_analysis.xlsx (from existing foci CSVs)"
        )
        self.coloc_calc_all_pairs.setChecked(True)
        coloc_outer.addWidget(self.coloc_calc_all_pairs)
        coloc_group.setLayout(coloc_outer)
        outer_layout.addWidget(coloc_group)

        # Wire FI UI state
        self.analyze_fi.toggled.connect(self._update_fi_ui_state)
        self.fi_mode_one_file.toggled.connect(self._update_fi_ui_state)
        self.fi_mode_match_by_name.toggled.connect(self._update_fi_ui_state)
        self.analyze_fi_contour.toggled.connect(self._update_fi_contour_state)
        self.file_path.textChanged.connect(self._update_fi_feedback)
        self._update_fi_ui_state()
        return outer

    def _select_fi_file(self) -> None:
        from PyQt6.QtWidgets import QFileDialog

        file_path, _ = QFileDialog.getOpenFileName(
            self, "Select FI Channel TIFF File", get_default_browse_dir(), "TIFF Files (*.tif *.tiff);;All Files (*)"
        )
        if file_path:
            self._single_fi_path = Path(file_path)
            self._update_fi_feedback()

    def _is_match_by_name_mode(self) -> bool:
        return bool(getattr(self, "fi_mode_match_by_name", None) and self.fi_mode_match_by_name.isChecked())

    def _update_fi_contour_state(self, *_args) -> None:
        try:
            self.fi_contour_thickness.setEnabled(bool(self.analyze_fi.isChecked() and self.analyze_fi_contour.isChecked()))
        except Exception:
            pass

    def _update_fi_ui_state(self, _checked=None) -> None:
        analyze_fi_on = bool(getattr(self, "analyze_fi", None) and self.analyze_fi.isChecked())
        try:
            self.fi_select_btn.setVisible(analyze_fi_on and not self._is_match_by_name_mode())
        except Exception:
            pass
        try:
            self.fi_feedback.setVisible(analyze_fi_on)
        except Exception:
            pass
        self._update_fi_contour_state()
        self._update_fi_feedback()

    def _update_fi_feedback(self) -> None:
        if not getattr(self, "fi_feedback", None):
            return
        if not self.analyze_fi.isChecked():
            self.fi_feedback.setPlainText("Analyze FI is off. Enable it to see FI pairings.")
            return
        if self._is_match_by_name_mode():
            masks = self._get_mask_files_for_feedback()
            if not masks:
                self.fi_feedback.setPlainText("Set mask path (file or folder) to see FI pairings.")
                return
            fi_lines = []
            total = 0
            for mask_path in masks:
                fi_paths = self._resolve_fi_paths_for_mask(mask_path)
                total += len(fi_paths)
                fi_lines.append(f"• {mask_path.name} → {', '.join(p.name for p in fi_paths) if fi_paths else '(no FI found)'}")
            header = f"{len(masks)} mask file(s) → {total} analyses\n\n"
            self.fi_feedback.setPlainText(header + "Analyze FI (mask → FI channel(s)):\n" + "\n".join(fi_lines))
        else:
            if self._single_fi_path and self._single_fi_path.exists():
                self.fi_feedback.setPlainText(f"FI file: {self._single_fi_path.name} will be used for all masks.")
            else:
                self.fi_feedback.setPlainText("Select one FI file (button above) to use for all masks.")

    # ---- analysis runner (ported; foci detection removed) -------------
    def analyze_masks(self) -> None:
        if not self.file_path.text():
            QMessageBox.warning(self, "No Selection", "Please select a mask file or folder first.")
            return
        if not Path(self.file_path.text()).exists():
            QMessageBox.warning(self, "Not Found", "The selected file or folder does not exist.")
            return

        # Run the package script as a module so we always use bactimate.scripts.analyze_cell_areas
        self._analyze_script_module = "bactimate.scripts.analyze_cell_areas"
        self._post_analysis_coloc_masks: list[Path] | None = None

        is_folder = self._is_folder_mode()
        if is_folder:
            # Collect per-task failures so folder runs don't appear "silent".
            self._folder_failures: list[tuple[str, str]] = []
            folder = Path(self.file_path.text()).resolve()
            use_filtered = bool(getattr(self, "use_filtered_masks_cb", None) and self.use_filtered_masks_cb.isChecked())
            mask_files_set = set()
            if use_filtered:
                mask_files_set.update(folder.glob("*_mask_filtered.tif"))
                mask_files_set.update(folder.glob("*_mask_filtered.tiff"))
            if not mask_files_set:
                mask_files_set.update(folder.glob("*_masks.tif"))
                mask_files_set.update(folder.glob("*_masks.tiff"))
            mask_files = sorted(mask_files_set, key=lambda x: x.name.lower())
            if not mask_files:
                QMessageBox.warning(self, "No Mask Files Found", "No mask files found in the selected folder.")
                return

            selected_files = show_file_selection_dialog(
                self,
                mask_files,
                title="Select Mask Files for Analysis",
                message=f"Select mask files to analyze (all {len(mask_files)} files selected by default):",
            )
            if selected_files is None or len(selected_files) == 0:
                return

            if self.analyze_fi.isChecked() and not self._is_match_by_name_mode() and not (
                getattr(self, "_single_fi_path", None) and self._single_fi_path.exists()
            ):
                QMessageBox.warning(
                    self,
                    "FI File Required",
                    "Analyze FI is on with 'One FI file for all masks' but no file is selected.",
                )
                return

            self.analyze_btn.setEnabled(False)
            self.analyze_btn.setText(f"Analyzing... (0/{len(selected_files)})")

            self.folder_files = selected_files
            self.folder_results = []
            self.folder_output_paths = []
            batch_fi_on = self.analyze_fi.isChecked() and self._is_match_by_name_mode()
            single_fi_path = getattr(self, "_single_fi_path", None) if (self.analyze_fi.isChecked() and not batch_fi_on) else None
            if single_fi_path and not single_fi_path.exists():
                single_fi_path = None
            self.folder_fi_tasks: list[tuple[Path, Path | None, Path]] = []
            self.folder_phase_per_mask: dict[Path, Path | None] = {}

            def _add_task(mask_p: Path, fi_p: Path | None, out_p: Path) -> None:
                self.folder_fi_tasks.append((mask_p, fi_p, out_p))

            for mask_path in selected_files:
                mask_path = mask_path.resolve()
                base_dir = base_dir_for_mask(mask_path)
                stem = stem_from_mask_path(mask_path)
                analysis_dir = analysis_dir_from_base(base_dir)
                if batch_fi_on:
                    # Primary: explicit FI stacks (name-matched).
                    # Fallback: if none exist, treat foci stacks as FI inputs (common datasets only have *_foci_*.tif).
                    fi_paths = self._resolve_fi_paths_for_mask(mask_path)
                    if not fi_paths:
                        fi_paths = self._resolve_foci_paths_for_mask(mask_path)
                    if not fi_paths:
                        _add_task(mask_path, None, analysis_dir / f"{stem}_analysis.xlsx")
                    else:
                        for fi_path in fi_paths:
                            suffix = self._fi_output_suffix(fi_path, stem)
                            _add_task(mask_path, fi_path, analysis_dir / f"{stem}_analysis{suffix}.xlsx")
                else:
                    _add_task(mask_path, single_fi_path, analysis_dir / f"{stem}_analysis.xlsx")

            # Validate
            validation_errors = []
            for mask_path, fi_path, _out in self.folder_fi_tasks:
                ok, _missing = paths_exist([mask_path])
                if not ok:
                    validation_errors.append(f"{mask_path.name}: mask file not found.")
                    continue
                if fi_path is not None:
                    ok, _missing = paths_exist([fi_path])
                    if not ok:
                        validation_errors.append(f"{mask_path.name}: FI file not found.")
                        continue
                    ok, err = mask_fi_spatial_shape_match(mask_path, fi_path)
                    if not ok and err:
                        validation_errors.append(f"{mask_path.name}: {err}")
            if validation_errors:
                QMessageBox.warning(self, "Validation failed", "Cannot start analysis:\n\n" + "\n".join(validation_errors[:10]))
                self.analyze_btn.setEnabled(True)
                self.analyze_btn.setText(self._primary_button_label())
                return

            self.current_task_index = 0
            self._process_next_folder_file()
            return

        # Single file mode: run one task
        input_file = Path(self.file_path.text()).resolve()
        if getattr(self, "use_filtered_masks_cb", None) is not None and self.use_filtered_masks_cb.isChecked():
            try:
                name_lower = input_file.name.lower()
                if name_lower.endswith(("_masks.tif", "_masks.tiff")) and "_mask_filtered" not in name_lower:
                    candidate = input_file.parent / (stem_from_mask_path(input_file) + "_mask_filtered.tif")
                    if candidate.exists():
                        input_file = candidate
            except Exception:
                pass

        mask_path = input_file
        base_dir = base_dir_for_mask(mask_path)
        stem = stem_from_mask_path(mask_path)
        analysis_dir = analysis_dir_from_base(base_dir)
        out = analysis_dir / f"{stem}_analysis.xlsx"

        # If Analyze FI is on and we're in match-by-name mode, run one task per channel
        # (FI stacks if present, otherwise foci stacks), then merge back into {stem}_analysis.xlsx.
        if self.analyze_fi.isChecked() and self._is_match_by_name_mode():
            fi_paths = self._resolve_fi_paths_for_mask(mask_path)
            if not fi_paths:
                fi_paths = self._resolve_foci_paths_for_mask(mask_path)

            self.folder_files = [mask_path]
            self.folder_results = []
            self.folder_output_paths = []
            self.folder_fi_tasks = []

            if not fi_paths:
                self.folder_fi_tasks.append((mask_path, None, out))
            else:
                for fi_path in fi_paths:
                    suffix = self._fi_output_suffix(fi_path, stem)
                    self.folder_fi_tasks.append((mask_path, fi_path, analysis_dir / f"{stem}_analysis{suffix}.xlsx"))

            self.analyze_btn.setEnabled(False)
            self.analyze_btn.setText(f"Analyzing... (0/{len(self.folder_fi_tasks)})")
            self.current_task_index = 0
            self._process_next_folder_file()
            return

        fi_path = None
        if self.analyze_fi.isChecked() and not self._is_match_by_name_mode():
            fi_path = getattr(self, "_single_fi_path", None)
        self._run_one_analysis_task(mask_path, fi_path, out)

    def _run_one_analysis_task(self, mask_path: Path, fi_path: Path | None, out_path: Path) -> None:
        ok, _missing = paths_exist([mask_path] + ([fi_path] if fi_path is not None else []))
        if not ok:
            QMessageBox.warning(self, "Not Found", "One or more inputs do not exist.")
            return
        if fi_path is not None:
            ok, err = mask_fi_spatial_shape_match(mask_path, fi_path)
            if not ok and err:
                QMessageBox.warning(self, "Validation failed", err)
                return

        cmd_args = self._build_analyze_cell_areas_args(mask_path=mask_path, out_path=out_path, fi_path=fi_path)

        def _done(ok: bool, msg: str) -> None:
            # Verify the expected artifact exists when the script reported success.
            try:
                if ok and out_path is not None and isinstance(out_path, Path) and not out_path.exists():
                    ok = False
                    msg = (msg or "").strip()
                    extra = f"Expected output was not created:\n{out_path}"
                    msg = (msg + "\n\n" + extra).strip() if msg else extra
            except Exception:
                pass
            if hasattr(self.parent_window, "log_output") and msg:
                self.parent_window.log_output(msg + ("\n" if not msg.endswith("\n") else ""))
            QMessageBox.information(
                self,
                "Analysis complete" if ok else "Analysis failed",
                msg or ("Done" if ok else "Failed"),
            )

        # Use existing WorkerThread subprocess runner (expects script_path + cmd_args)
        # We run analyze_cell_areas as a module: script_path="-m", cmd_args=[module, ...args]
        t = WorkerThread("-m", cmd_args, parent=self)
        t.finished.connect(_done)
        if hasattr(t, "progress") and hasattr(self.parent_window, "log_output"):
            t.progress.connect(lambda s: self.parent_window.log_output(str(s)))
        t.start()

    def _process_next_folder_file(self) -> None:
        tasks = getattr(self, "folder_fi_tasks", []) or []
        idx = int(getattr(self, "current_task_index", 0))
        if idx >= len(tasks):
            # Done; optionally merge FI channels
            try:
                self._merge_fi_channels_into_one_excel()
            except Exception:
                pass
            # Post-analysis colocalization (CSV-driven, no re-detection)
            try:
                if hasattr(self, "coloc_calc_all_pairs") and self.coloc_calc_all_pairs.isChecked():
                    from bactimate.gui.tabs._analyze_masks_colocalization import run_post_analysis_colocalization_for_masks

                    # Determine which masks succeeded (all tasks for that mask).
                    results = getattr(self, "folder_results", []) or []
                    ok_masks: list[Path] = []
                    by_mask: dict[Path, list[bool]] = {}
                    for (m, _fi, _out), ok in zip(tasks, results, strict=False):
                        by_mask.setdefault(Path(m), []).append(bool(ok))
                    for m, oks in by_mask.items():
                        if oks and all(oks):
                            ok_masks.append(Path(m))
                    if ok_masks:
                        run_post_analysis_colocalization_for_masks(self, ok_masks)
            except Exception:
                pass

            # Surface failures to the user (folder runs otherwise only log).
            try:
                failures = getattr(self, "_folder_failures", []) or []
                if failures:
                    head = "\n\n".join([f"- {name}\n{msg}".strip() for name, msg in failures[:3]])
                    tail = ""
                    if len(failures) > 3:
                        tail = f"\n\n(+{len(failures) - 3} more; see log)"
                    QMessageBox.warning(
                        self,
                        "Analysis finished with errors",
                        "One or more analysis tasks failed or produced no output.\n\n"
                        + head
                        + tail,
                    )
            except Exception:
                pass
            self.analyze_btn.setEnabled(True)
            self.analyze_btn.setText(self._primary_button_label())
            return
        mask_path, fi_path, out_path = tasks[idx]
        self.analyze_btn.setText(f"Analyzing... ({idx+1}/{len(tasks)})")

        cmd_args = self._build_analyze_cell_areas_args(mask_path=mask_path, out_path=out_path, fi_path=fi_path)
        th = WorkerThread("-m", cmd_args, parent=self)

        def _on_done(success: bool, message: str):
            # Verify the expected output exists; treat missing artifact as failure.
            try:
                if success and out_path is not None and not Path(out_path).exists():
                    success = False
                    extra = f"Expected output was not created: {out_path}"
                    message = (str(message or "").strip() + "\n" + extra).strip()
            except Exception:
                pass
            getattr(self, "folder_results", []).append(bool(success))
            getattr(self, "folder_output_paths", []).append(str(out_path))
            if hasattr(self.parent_window, "log_output") and message:
                self.parent_window.log_output(message + ("\n" if not message.endswith("\n") else ""))
            if not success:
                try:
                    lst = getattr(self, "_folder_failures", None)
                    if isinstance(lst, list):
                        lst.append((Path(mask_path).name, str(message or "Failed").strip()))
                except Exception:
                    pass
            self.current_task_index = idx + 1
            self._process_next_folder_file()

        th.finished.connect(_on_done)
        if hasattr(th, "progress") and hasattr(self.parent_window, "log_output"):
            th.progress.connect(lambda s: self.parent_window.log_output(str(s)))
        th.start()

    def _build_analyze_cell_areas_args(self, *, mask_path: Path, out_path: Path, fi_path: Path | None) -> list[str]:
        """
        Build CLI args for `bactimate.scripts.analyze_cell_areas` matching its argparse contract.

        Script signature:
          analyze_cell_areas masks --output OUT [--pixel-size ...] [--fi-image ...] ...
        """
        args: list[str] = [self._analyze_script_module, str(mask_path), "--output", str(out_path)]

        def _f(w, default: float) -> float:
            try:
                return float(w.value())
            except Exception:
                return float(default)

        def _opt_um2(w) -> float | None:
            try:
                v = float(w.value())
                return v if v > 0 else None
            except Exception:
                return None

        # Core geometry / filtering parameters
        args += ["--pixel-size", str(_f(getattr(self, "pixel_size", None), 0.108))]
        # Morphology/size filtering is handled upstream in the Filtering stage.

        # Output detail toggles (script uses positive flags)
        detailed_on = bool(getattr(self, "detailed", None) is not None and self.detailed.isChecked())
        # Post-analysis colocalization *requires* Per-Cell Data to exist in the workbook.
        # If the user enabled colocalization sheets, force --detailed so the contract holds.
        if getattr(self, "coloc_calc_all_pairs", None) is not None and self.coloc_calc_all_pairs.isChecked():
            detailed_on = True
        if detailed_on:
            args += ["--detailed"]
        if getattr(self, "create_vis", None) is not None and self.create_vis.isChecked():
            args += ["--visualization"]
        # NOTE: We intentionally do not emit --export-csv from the GUI.
        # The CSV sidecars duplicate the content already present in the Excel workbook
        # and tend to clutter `_analysis/` for end users.
        # Note: analysis intentionally ignores exclude-border; batch filtering handles it.

        # FI options
        if getattr(self, "analyze_fi", None) is not None and self.analyze_fi.isChecked() and fi_path is not None:
            args += ["--fi-image", str(fi_path)]
            if getattr(self, "analyze_fi_contour", None) is not None and self.analyze_fi_contour.isChecked():
                args += ["--fi-contour-thickness", str(int(self.fi_contour_thickness.value()))]
            else:
                args += ["--no-fi-contour"]

        return args

    # ---- helpers needed by merge/summary -----------------------------
    def _fi_output_suffix(self, fi_path: Path | None, stem: str) -> str:
        if fi_path is None:
            return ""
        try:
            tag = self._extract_channel_tag_for_mask_stem(fi_path.stem, stem) or ""
        except Exception:
            tag = ""
        if tag:
            return "_" + str(tag)
        # fallback: use remainder after stem_
        s = fi_path.stem
        if s.startswith(stem + "_"):
            rem = s[len(stem) + 1 :]
            if rem:
                return "_" + rem
        return ""

    # ---- combined summary + FI merge (ported verbatim-ish) -----------
    def _match_log_from_folder_tasks(self) -> tuple[int, int, list[str]]:
        tasks = getattr(self, "folder_fi_tasks", [])
        results = getattr(self, "folder_results", [])
        if len(results) != len(tasks):
            return 0, 0, []
        seen_stem: set[str] = set()
        lines: list[str] = []
        mask_stems: set[str] = set()
        n_tif = 0
        for i, (mask_file, fi_path, _out) in enumerate(tasks):
            if not results[i]:
                continue
            stem = mask_file.stem.replace("_masks", "")
            tif_name = fi_path.name if fi_path is not None else "—"
            if fi_path is not None:
                n_tif += 1
            mask_stems.add(stem)
            if stem not in seen_stem:
                seen_stem.add(stem)
                lines.append(f"{stem}_masks -> {tif_name}")
            else:
                lines.append(f"{stem}_masks with {tif_name}")
        return len(mask_stems), n_tif, lines

    def _build_combined_summary_excel(self, excel_paths, output_path, match_log=None, log_callback=None):
        # Copied from legacy implementation (see original `analyze_masks.py`)
        def _log(msg: str) -> None:
            if log_callback:
                log_callback(msg)

        import re
        import pandas as pd
        from openpyxl import load_workbook
        from openpyxl.utils.dataframe import dataframe_to_rows

        def _sanitize_sheet_name(name: str, max_len: int = 31) -> str:
            s = re.sub(r'[\\/:*?\[\]]', '_', str(name))
            return s[:max_len] if len(s) > max_len else s

        if not excel_paths:
            _log("Combined summary: no paths provided.\n")
            return False
        frames_long = []
        missing = []
        no_summary = []
        per_cell_cache: dict[Path, pd.DataFrame | None] = {}
        for p in excel_paths:
            p = Path(p)
            if not p.exists():
                missing.append(p.name)
                continue
            try:
                with pd.ExcelFile(p) as xl:
                    df = pd.read_excel(xl, sheet_name="Summary")
                    per_cell = None
                    try:
                        per_cell = pd.read_excel(xl, sheet_name="Per-Cell Data")
                    except Exception:
                        pass
                    per_cell_cache[p] = per_cell
            except Exception as e:
                no_summary.append((p.name, str(e)))
                continue
            if df is None or df.empty:
                no_summary.append((p.name, "empty Summary sheet"))
                continue
            df = df.copy()
            df["File"] = p.stem.replace("_analysis", "").replace("_masks_analysis", "")
            frames_long.append(df)
        if missing:
            _log(
                f"Combined summary: {len(missing)} path(s) not found: {', '.join(missing[:5])}{' ...' if len(missing) > 5 else ''}\n"
            )
        if no_summary and not frames_long:
            _log(f"Combined summary: no Summary sheet could be read. First failure: {no_summary[0][0]} — {no_summary[0][1]}\n")
        if not frames_long:
            return False
        long_df = pd.concat(frames_long, ignore_index=True)
        cols = [c for c in ["File", "Frame"] if c in long_df.columns]
        cols += [c for c in long_df.columns if c not in cols]
        long_df = long_df[cols]
        wide_df = None
        wide_err = None
        try:
            if "Frame" not in long_df.columns:
                wide_err = "Summary sheets have no 'Frame' column"
            else:
                frames = sorted(long_df["Frame"].dropna().unique().astype(int).tolist())
                if not frames:
                    wide_err = "Summary sheets have no Frame values"
                else:
                    wide_data = {"Frame": frames}
                    for f in long_df["File"].unique().tolist():
                        sub = long_df[long_df["File"] == f].drop_duplicates(subset=["Frame"], keep="first").set_index("Frame")
                        for col in ["Total Cell Area (µm²)", "Total Cell Area (Relative)", "Number of Cells"]:
                            if col in sub.columns:
                                wide_data[f"{f}_{col}"] = sub.reindex(frames)[col].values
                    wide_df = pd.DataFrame(wide_data)
        except Exception as e:
            wide_err = str(e)
            wide_df = None
        if wide_df is None or wide_df.empty:
            if wide_err:
                _log(f"Combined summary: wide table build failed — {wide_err}\n")
            return False
        out_path = Path(output_path)
        existing_log_lines: list[str] = []
        if out_path.exists() and match_log is not None:
            try:
                wb_existing = load_workbook(output_path, data_only=True)
                sheet_name_wide = "Combined Summary (wide)"
                if sheet_name_wide in wb_existing.sheetnames:
                    ws_existing = wb_existing[sheet_name_wide]
                    for r in range(2, min(ws_existing.max_row + 1, 5000)):
                        val = ws_existing.cell(row=r, column=7).value
                        if val is None or not str(val).strip():
                            continue
                        s = str(val).strip()
                        if " -> " in s or " with " in s:
                            existing_log_lines.append(s)
                wb_existing.close()
            except Exception:
                pass
        try:
            from openpyxl import Workbook

            wb = Workbook()
            ws = wb.active
            ws.title = "Combined Summary (wide)"
            for row in dataframe_to_rows(wide_df, index=False, header=True):
                ws.append(row)
            if match_log is not None:
                n_masks, n_tif, lines = match_log
                ws["G1"] = "Logs"
                ws["G2"] = f"{n_masks} mask(s) and {n_tif} tif file(s) matched"
                rowi = 3
                for line in existing_log_lines:
                    ws.cell(row=rowi, column=7, value=line)
                    rowi += 1
                for line in lines:
                    ws.cell(row=rowi, column=7, value=line)
                    rowi += 1
            seen_sheet_names = set(wb.sheetnames)
            for p in excel_paths:
                p = Path(p)
                if not p.exists():
                    continue
                per_cell = per_cell_cache.get(p)
                if per_cell is None or per_cell.empty:
                    continue
                base_name = p.stem.replace("_analysis", "").replace("_masks_analysis", "")
                sheet_name = _sanitize_sheet_name(base_name)
                if sheet_name in seen_sheet_names:
                    stem2 = sheet_name[:28] if len(sheet_name) > 28 else sheet_name
                    n = 1
                    while f"{stem2}_{n}" in seen_sheet_names:
                        n += 1
                    sheet_name = f"{stem2}_{n}"
                seen_sheet_names.add(sheet_name)
                new_ws = wb.create_sheet(title=sheet_name)
                for row in dataframe_to_rows(per_cell, index=False, header=True):
                    new_ws.append(row)
            wb.save(output_path)
        except Exception as e:
            _log(f"Combined summary: write failed — {e}\n")
            return False
        return True

    def _merge_fi_channels_into_one_excel(self) -> None:
        # Copied from legacy implementation with minimal changes
        import pandas as pd

        tasks = getattr(self, "folder_fi_tasks", [])
        results = getattr(self, "folder_results", [])
        if len(results) != len(tasks):
            return

        def _col_key(raw: str) -> str:
            txt = str(raw or "").strip()
            if not txt:
                return "FI"
            out = []
            for ch in txt:
                out.append(ch if ch.isalnum() or ch == "_" else "_")
            key = "".join(out).strip("_")
            while "__" in key:
                key = key.replace("__", "_")
            return key or "FI"

        by_stem: dict[str, list[tuple[Path, Path | None, str]]] = {}
        for i, (mask_file, fi_path, output_path) in enumerate(tasks):
            if not results[i] or not output_path.exists():
                continue
            stem = mask_file.stem.replace("_masks", "")
            by_stem.setdefault(stem, [])
            channel_label = "FI"
            if fi_path is not None:
                suffix = self._fi_output_suffix(fi_path, stem)
                channel_label = (suffix[1:] if suffix.startswith("_") else "FI") or "FI"
            by_stem[stem].append((output_path, fi_path, channel_label))

        for stem, group in by_stem.items():
            phase_entries = [(p, fp, lbl) for p, fp, lbl in group if fp is None]
            fi_entries = [(p, fp, lbl) for p, fp, lbl in group if fp is not None]
            if not fi_entries:
                continue
            analysis_dir = group[0][0].parent
            merged_path = analysis_dir / f"{stem}_analysis.xlsx"

            try:
                summary_df = None
                merged = None
                first_excel = phase_entries[0][0] if phase_entries else fi_entries[0][0]

                for _p, _fp, _lbl in fi_entries:
                    try:
                        summary_df = pd.read_excel(_p, sheet_name="Summary")
                        break
                    except Exception:
                        continue
                if summary_df is None and phase_entries:
                    try:
                        summary_df = pd.read_excel(phase_entries[0][0], sheet_name="Summary")
                    except Exception:
                        pass

                try:
                    head = pd.read_excel(first_excel, sheet_name="Per-Cell Data", nrows=0)
                    key_cols = ["Frame", "Cell_ID"] if "Cell_ID" in head.columns else ["Frame", "Cell Label"]
                except Exception:
                    key_cols = ["Frame", "Cell Label"]

                fi_col_names = [
                    "Mean FI Intensity",
                    "Mean FI Intensity (whole cell)",
                    "Contour FI Intensity",
                    "FI Ratio (Contour/Mean)",
                ]
                fi_col_bases = list(fi_col_names)

                def _strip_bracket_tag(col: str) -> tuple[str, str | None]:
                    """
                    Parse columns of the form "Base [TAG]" -> ("Base", "TAG").
                    Returns (original, None) when no bracket tag is present.
                    """
                    s = str(col or "")
                    if "[" not in s or "]" not in s:
                        return s, None
                    # Only treat trailing bracket blocks as tags.
                    i = s.rfind("[")
                    j = s.rfind("]")
                    if i < 0 or j < 0 or j < i:
                        return s, None
                    base = s[:i].rstrip()
                    tag = s[i + 1 : j].strip()
                    if not base or not tag:
                        return s, None
                    if j != (len(s) - 1):
                        return s, None
                    return base, tag

                def _canonicalize_fi_columns(df: "pd.DataFrame") -> tuple["pd.DataFrame", dict[str, str]]:
                    """
                    Return (df2, colmap) where df2 has FI columns renamed from e.g.
                    "Contour FI Intensity [GFP]" -> "Contour FI Intensity" (base name only),
                    so downstream merging can consistently suffix with (channel_label).
                    """
                    if df is None or df.empty:
                        return df, {}
                    colmap: dict[str, str] = {}
                    for c in list(df.columns):
                        if not isinstance(c, str):
                            continue
                        base, _tag = _strip_bracket_tag(c)
                        if base in fi_col_bases and base != c:
                            colmap[c] = base
                    if not colmap:
                        return df, {}
                    return df.rename(columns=colmap), colmap

                if phase_entries:
                    phase_path = phase_entries[0][0]
                    try:
                        df_phase = pd.read_excel(phase_path, sheet_name="Per-Cell Data")
                        if df_phase is not None and not df_phase.empty:
                            merged = df_phase.copy()
                    except Exception:
                        pass

                for excel_path, _fp, channel_label in fi_entries:
                    try:
                        df = pd.read_excel(excel_path, sheet_name="Per-Cell Data")
                    except Exception:
                        df = None
                    if df is None or df.empty:
                        continue

                    # Normalize any channel-tagged FI columns (e.g. "... [GFP]") back to base names.
                    df, _fi_colmap = _canonicalize_fi_columns(df)

                    has_fi = [c for c in fi_col_names if c in df.columns]
                    rename = {c: f"{c} ({channel_label})" for c in has_fi}
                    sub = df.rename(columns=rename)
                    merge_on = [c for c in key_cols if merged is not None and c in merged.columns and c in sub.columns]

                    def is_fi_or_foci_col(c: str) -> bool:
                        if c in key_cols:
                            return False
                        up = str(c).upper()
                        if "FI" in str(c) or "FOCI" in up or c == "Number of Foci":
                            return True
                        if str(c).startswith("foci_ids_") or str(c).endswith("_foci_count"):
                            return True
                        return False

                    if merged is None:
                        merged = sub.copy()
                    else:
                        fi_new = [c for c in sub.columns if c not in merge_on and is_fi_or_foci_col(c) and c not in merged.columns]
                        if merge_on and fi_new:
                            merged = merged.merge(sub[merge_on + fi_new], on=merge_on, how="outer")

                if summary_df is None:
                    continue

                # ------------------------------------------------------------------
                # Ensure merged Summary contains per-channel FI frame aggregates when FI is present.
                # The underlying analysis script writes FI at the per-cell level; for multi-channel
                # batch runs we compute frame-mean FI metrics and append them to the merged Summary.
                # ------------------------------------------------------------------
                try:
                    if merged is not None and not merged.empty and "Frame" in merged.columns:
                        summary_df = summary_df.copy()
                        used_cols = set(str(c) for c in summary_df.columns)
                        frame_series = pd.to_numeric(merged["Frame"], errors="coerce")
                        merged_with_frame = merged.copy()
                        merged_with_frame["_frame_int"] = frame_series.astype("Int64")

                        def _add_frame_mean(col_name: str, out_col: str) -> None:
                            nonlocal summary_df
                            if out_col in used_cols:
                                return
                            if col_name not in merged_with_frame.columns:
                                return
                            s = pd.to_numeric(merged_with_frame[col_name], errors="coerce")
                            tmp = pd.DataFrame({"Frame": merged_with_frame["_frame_int"], "__v": s})
                            tmp = tmp.dropna(subset=["Frame"])
                            if tmp.empty:
                                return
                            means = tmp.groupby("Frame", dropna=True)["__v"].mean()
                            # Summary 'Frame' is 1-based int; map means back by Frame.
                            try:
                                sf = pd.to_numeric(summary_df["Frame"], errors="coerce").astype("Int64")
                            except Exception:
                                return
                            summary_df[out_col] = sf.map(lambda f: float(means.get(int(f), np.nan)) if f is not None and pd.notna(f) else np.nan)
                            used_cols.add(out_col)

                        import numpy as np  # local to keep function standalone

                        for _excel_path, _fp, ch in fi_entries:
                            # These names are created earlier: "<base> (<channel_label>)"
                            for base in fi_col_bases:
                                src = f"{base} ({ch})"
                                # Put same-named columns into Summary as frame means.
                                _add_frame_mean(src, src)
                except Exception:
                    pass

                # Simplify foci columns
                if merged is not None and not merged.empty:
                    simple_foci_cols: set[str] = set()
                    for col in list(merged.columns):
                        if not isinstance(col, str):
                            continue
                        if col.startswith("foci_ids_"):
                            ch = col[len("foci_ids_") :].strip() or "FI"
                            alias = f"foci id [{ch}]"
                            merged[alias] = merged[col]
                            simple_foci_cols.add(alias)
                        elif col.endswith("_foci_count"):
                            ch = col[: -len("_foci_count")].strip() or "FI"
                            alias = f"num foci [{ch}]"
                            merged[alias] = merged[col]
                            simple_foci_cols.add(alias)

                    def _is_foci_clutter_col(c: str) -> bool:
                        if c in simple_foci_cols:
                            return False
                        if c.startswith("num_coloc_foci_"):
                            return False
                        up = c.upper()
                        if c.startswith("foci_ids_") or c.endswith("_foci_count"):
                            return True
                        if "FOCI" in up:
                            return True
                        return False

                    drop_cols = [c for c in list(merged.columns) if isinstance(c, str) and _is_foci_clutter_col(c)]
                    if drop_cols:
                        merged = merged.drop(columns=drop_cols, errors="ignore")

                # Write merged workbook and per-channel foci sheets (if present)
                foci_by_channel: dict[str, list[pd.DataFrame]] = {}
                for excel_path, _fp, ch_lbl in fi_entries:
                    try:
                        df_foci = pd.read_excel(excel_path, sheet_name="Foci")
                    except Exception:
                        df_foci = None
                    if df_foci is None or df_foci.empty:
                        continue
                    df_foci = df_foci.copy()
                    if "channel_name" not in df_foci.columns:
                        df_foci["channel_name"] = ch_lbl
                    foci_by_channel.setdefault(ch_lbl, []).append(df_foci)

                with pd.ExcelWriter(merged_path, engine="openpyxl") as writer:
                    summary_df.to_excel(writer, sheet_name="Summary", index=False)
                    if merged is not None and not merged.empty:
                        merged.to_excel(writer, sheet_name="Per-Cell Data", index=False)
                    for ch in sorted(foci_by_channel.keys(), key=lambda s: str(s).lower()):
                        parts = foci_by_channel.get(ch) or []
                        if not parts:
                            continue
                        ch_df = pd.concat(parts, ignore_index=True, sort=False)
                        ch_df.to_excel(writer, sheet_name=f"Foci [{ch}]", index=False)
            except Exception:
                get_logger().exception("Merge FI channels failed")
                continue

    # ---- analyze session params extension ----------------------------
    def get_session_params(self) -> dict[str, Any]:
        d = super().get_session_params()
        d.update(
            {
                "analyze_fi": bool(self.analyze_fi.isChecked()) if hasattr(self, "analyze_fi") else False,
                "fi_match_by_name": bool(self.fi_mode_match_by_name.isChecked()) if hasattr(self, "fi_mode_match_by_name") else True,
                "single_fi_path": str(self._single_fi_path) if getattr(self, "_single_fi_path", None) is not None else "",
                "analyze_fi_contour": bool(self.analyze_fi_contour.isChecked()) if hasattr(self, "analyze_fi_contour") else False,
                "fi_contour_thickness": int(self.fi_contour_thickness.value()) if hasattr(self, "fi_contour_thickness") else 1,
                "use_filtered_masks": bool(self.use_filtered_masks_cb.isChecked()) if hasattr(self, "use_filtered_masks_cb") else True,
                "coloc_merge_max_distance": float(self.coloc_merge_max_distance.value()) if hasattr(self, "coloc_merge_max_distance") else 8.0,
                "coloc_merge_threshold": float(self.coloc_merge_threshold.value()) if hasattr(self, "coloc_merge_threshold") else 4.0,
                "coloc_merge_same_cell_only": bool(self.coloc_merge_same_cell_only.isChecked()) if hasattr(self, "coloc_merge_same_cell_only") else True,
                "coloc_calc_all_pairs": bool(self.coloc_calc_all_pairs.isChecked()) if hasattr(self, "coloc_calc_all_pairs") else False,
            }
        )
        return d

    def set_session_params(self, data: dict[str, Any]) -> None:
        super().set_session_params(data)
        apply_to_widgets(
            data,
            [
                ("analyze_fi", self.analyze_fi),
                ("analyze_fi_contour", self.analyze_fi_contour),
                ("fi_contour_thickness", self.fi_contour_thickness),
                ("use_filtered_masks", self.use_filtered_masks_cb),
                ("coloc_merge_max_distance", self.coloc_merge_max_distance),
                ("coloc_merge_threshold", self.coloc_merge_threshold),
                ("coloc_merge_same_cell_only", self.coloc_merge_same_cell_only),
                ("coloc_calc_all_pairs", self.coloc_calc_all_pairs),
            ],
        )
        fi_match = data.get("fi_match_by_name", data.get("batch_fi", True))
        if hasattr(self, "fi_mode_match_by_name"):
            self.fi_mode_match_by_name.setChecked(bool(fi_match))
            self.fi_mode_one_file.setChecked(not bool(fi_match))
        if "single_fi_path" in data and data["single_fi_path"]:
            try:
                self._single_fi_path = Path(data["single_fi_path"])
            except Exception:
                self._single_fi_path = None
        self._update_fi_ui_state()

