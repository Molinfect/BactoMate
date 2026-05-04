from __future__ import annotations

import itertools
from collections import defaultdict
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any, Optional

import numpy as np

from bactimate.gui.analysis_paths import analysis_dir_from_base, base_dir_for_mask, stem_from_mask_path
from bactimate.gui.artifact_paths import (
    foci_csv_path,
    filtered_foci_csv_path,
    legacy_foci_csv_candidates,
    raw_foci_csv_path,
)
from bactimate.gui.utils import safe_imread


def _log_tab(tab, msg: str) -> None:
    try:
        pw = getattr(tab, "parent_window", None)
        if pw is not None and hasattr(pw, "log_output"):
            pw.log_output(msg if str(msg).endswith("\n") else (str(msg) + "\n"))
    except Exception:
        pass


def _safe_sheet_name(base: str, used: set[str]) -> str:
    bad = "\\/:*?[]"
    safe = "".join("_" if ch in bad else ch for ch in str(base)).strip() or "Sheet"
    safe = safe[:31]
    if safe not in used:
        used.add(safe)
        return safe
    stem_name = safe[:28] if len(safe) > 28 else safe
    n = 2
    cand = f"{stem_name}_{n}"
    while cand in used:
        n += 1
        cand = f"{stem_name}_{n}"
    used.add(cand)
    return cand


def _safe_col_key(raw: str) -> str:
    txt = str(raw or "").strip()
    if not txt:
        return "foci"
    out = []
    for ch in txt:
        out.append(ch if ch.isalnum() or ch == "_" else "_")
    key = "".join(out).strip("_")
    while "__" in key:
        key = key.replace("__", "_")
    return key or "foci"


def _excel_cell_value(v: Any) -> Any:
    """Convert arbitrary values into openpyxl-compatible cell scalars."""
    if v is None:
        return None
    try:
        if isinstance(v, np.generic):
            v = v.item()
    except Exception:
        pass
    if isinstance(v, (str, int, float, bool, date, datetime, time, Decimal)):
        return v
    if isinstance(v, (list, tuple, set, dict, np.ndarray)):
        return str(v)
    try:
        import pandas as pd

        if pd.isna(v):
            return None
    except Exception:
        pass
    return str(v)


def _load_mask_frames(mask_path: Path) -> dict[int, np.ndarray]:
    try:
        arr = np.asarray(safe_imread(str(mask_path)))
        if arr.ndim == 2:
            return {0: arr.astype(np.int32)}
        if arr.ndim >= 3:
            out: dict[int, np.ndarray] = {}
            for i in range(int(arr.shape[0])):
                out[i] = np.asarray(arr[i], dtype=np.int32)
            return out
    except Exception:
        return {}
    return {}


def _infer_frame_shift(
    *,
    per_cell_frames: set[int],
    channel_rows: list[tuple[str, list[dict[str, Any]]]],
    pair_tables: list[tuple[str, str, Any, int]],
) -> int:
    """0 if Summary Frame matches CSV/assigned frame; 1 if Summary is 1-based and CSV is 0-based."""
    import pandas as pd

    if not per_cell_frames:
        return 0
    observed_frames: set[int] = set()
    for _lbl, rows in channel_rows or []:
        for r in rows or []:
            try:
                observed_frames.add(int(r.get("frame", 0) or 0))
            except Exception:
                continue
    for _la, _lb, df, _count in pair_tables or []:
        try:
            if df is None or df.empty or "frame" not in df.columns:
                continue
            for v in df["frame"].tolist():
                if v is None or pd.isna(v):
                    continue
                observed_frames.add(int(v))
        except Exception:
            continue
    if not observed_frames:
        return 0
    try:
        overlap_raw = len(per_cell_frames.intersection(observed_frames))
        overlap_plus1 = len(per_cell_frames.intersection({f + 1 for f in observed_frames}))
        if overlap_plus1 > overlap_raw:
            return 1
    except Exception:
        return 0
    return 0


_SUMMARY_COL_PREFIXES = (
    "Total Foci Count (",
    "Total Foci Count within Cells (",
    "Valid Pairs Count (",
)


def _strip_managed_summary_columns(df: "pd.DataFrame") -> "pd.DataFrame":
    """Remove columns written by post-colocalization Summary updates so reruns do not leave stale values."""
    drop = [
        c
        for c in df.columns
        if isinstance(c, str) and any(c.startswith(p) for p in _SUMMARY_COL_PREFIXES)
    ]
    if drop:
        return df.drop(columns=drop, errors="ignore")
    return df


def _pair_column_name(la: str, lb: str) -> str:
    a, b = sorted((str(la), str(lb)), key=lambda s: s.lower())
    return f"Valid Pairs Count ({a}×{b})"


def _valid_pairs_counts_by_summary_frame(
    *,
    df: Any,
    frame_shift: int,
    same_cell_only: bool,
    rows_a: list[dict[str, Any]],
    rows_b: list[dict[str, Any]],
    summary_frames: set[int],
) -> dict[int, int]:
    """Map Summary Frame -> valid pair count for one channel pair."""
    import pandas as pd

    out: dict[int, int] = defaultdict(int)
    if df is None or df.empty or ("message" in df.columns and len(df.columns) == 1):
        return dict(out)
    if "cell_id_primary" not in df.columns:
        return dict(out)

    def _row_ok(row: Any) -> bool:
        try:
            if "is_valid" in df.columns and not bool(row.get("is_valid", True)):
                return False
            if same_cell_only and "same_cell" in df.columns and not bool(row.get("same_cell", False)):
                return False
            return True
        except Exception:
            return False

    has_frame_col = (
        "frame" in df.columns
        and len(df) > 0
        and bool(df["frame"].notna().any())
    )
    if has_frame_col:
        for _, row in df.iterrows():
            if not _row_ok(row):
                continue
            try:
                fr_i = int(row.get("frame", 0) or 0) + int(frame_shift)
            except Exception:
                continue
            out[fr_i] += 1
        return dict(out)

    # Single-frame analysis: pair rows have no 'frame' — attribute counts to Summary frame(s) from foci rows.
    try:
        cand = set()
        for r in rows_a or []:
            cand.add(int(r.get("frame", 0) or 0) + int(frame_shift))
        for r in rows_b or []:
            cand.add(int(r.get("frame", 0) or 0) + int(frame_shift))
        cand = {f for f in cand if f in summary_frames}
        if len(cand) == 1:
            only = next(iter(cand))
            total = sum(1 for _, row in df.iterrows() if _row_ok(row))
            out[only] = int(total)
            return dict(out)
    except Exception:
        pass
    total = sum(1 for _, row in df.iterrows() if _row_ok(row))
    if summary_frames:
        first = min(summary_frames)
        out[int(first)] = int(total)
    return dict(out)


def _persist_workbook_count_updates(
    *,
    analysis_path: Path,
    channel_rows: list[tuple[str, list[dict[str, Any]]]],
    raw_counts: dict[tuple[str, int], int],
    pair_tables: list[tuple[str, str, Any, int]],
    same_cell_only: bool,
) -> None:
    """Update Per-Cell Data foci/coloc counts and Summary per-frame foci / valid-pair columns (single write)."""
    import pandas as pd

    analysis_path = Path(analysis_path)
    if not analysis_path.exists():
        return

    try:
        per_cell = pd.read_excel(analysis_path, sheet_name="Per-Cell Data")
    except Exception:
        return
    if per_cell is None or per_cell.empty:
        return

    cell_col = "Cell_ID" if "Cell_ID" in per_cell.columns else ("Cell Label" if "Cell Label" in per_cell.columns else None)
    if cell_col is None:
        return
    has_frame = "Frame" in per_cell.columns

    try:
        summary = pd.read_excel(analysis_path, sheet_name="Summary")
    except Exception:
        summary = None

    per_cell_frames: set[int] = set()
    if has_frame:
        try:
            per_cell_frames = {
                int(v)
                for v in per_cell["Frame"].tolist()
                if v is not None and pd.notna(v)
            }
        except Exception:
            per_cell_frames = set()

    frame_shift = _infer_frame_shift(
        per_cell_frames=per_cell_frames,
        channel_rows=channel_rows,
        pair_tables=pair_tables,
    )

    # Reset existing count columns so reruns cannot leave stale values.
    existing_foci_cols = [c for c in per_cell.columns if str(c).endswith("_foci_count")]
    for c in existing_foci_cols:
        per_cell[c] = 0
    coloc_col = "coloc_pairs_count"
    if coloc_col not in per_cell.columns:
        per_cell[coloc_col] = 0
    else:
        per_cell[coloc_col] = 0

    # Foci per cell per channel.
    for label, rows in channel_rows or []:
        col = f"{_safe_col_key(label)}_foci_count"
        counts: dict[tuple[int, int], int] = {}
        for r in rows or []:
            try:
                if not bool(r.get("in_cell", False)):
                    continue
                cid = r.get("cell_id", None)
                if cid in (None, "", "None"):
                    continue
                cid_i = int(cid)
                if cid_i <= 0:
                    continue
                fr_i = int(r.get("frame", 0) or 0) + int(frame_shift)
                k = (fr_i, cid_i)
                counts[k] = counts.get(k, 0) + 1
            except Exception:
                continue
        if col not in per_cell.columns:
            per_cell[col] = 0
        if has_frame:
            per_cell[col] = [
                int(
                    counts.get(
                        (
                            int(fr) if fr is not None and pd.notna(fr) else 0,
                            int(cid) if cid is not None and pd.notna(cid) else -1,
                        ),
                        0,
                    )
                )
                for fr, cid in zip(per_cell["Frame"], per_cell[cell_col], strict=False)
            ]
        else:
            per_cell[col] = [
                int(counts.get((0, int(cid) if cid is not None and pd.notna(cid) else -1), 0))
                for cid in per_cell[cell_col]
            ]

    # Coloc pairs per cell: count valid pairs; optionally require same-cell.
    coloc_counts: dict[tuple[int, int], int] = {}
    for _la, _lb, df, _count in pair_tables or []:
        try:
            if df is None or df.empty or "message" in df.columns:
                continue
            if "cell_id_primary" not in df.columns:
                continue
            for _, row in df.iterrows():
                try:
                    if "is_valid" in df.columns and not bool(row.get("is_valid", True)):
                        continue
                    if same_cell_only and "same_cell" in df.columns and not bool(row.get("same_cell", False)):
                        continue
                    cid = row.get("cell_id_primary", None)
                    if cid in (None, "", "None"):
                        continue
                    cid_i = int(cid)
                    if cid_i <= 0:
                        continue
                    fr_i = 0
                    if "frame" in df.columns:
                        try:
                            fr_i = int(row.get("frame", 0) or 0) + int(frame_shift)
                        except Exception:
                            fr_i = int(frame_shift)
                    k = (fr_i, cid_i)
                    coloc_counts[k] = coloc_counts.get(k, 0) + 1
                except Exception:
                    continue
        except Exception:
            continue

    if has_frame:
        per_cell[coloc_col] = [
            int(
                coloc_counts.get(
                    (
                        int(fr) if fr is not None and pd.notna(fr) else 0,
                        int(cid) if cid is not None and pd.notna(cid) else -1,
                    ),
                    0,
                )
            )
            for fr, cid in zip(per_cell["Frame"], per_cell[cell_col], strict=False)
        ]
    else:
        per_cell[coloc_col] = [
            int(coloc_counts.get((0, int(cid) if cid is not None and pd.notna(cid) else -1), 0))
            for cid in per_cell[cell_col]
        ]

    # --- Summary sheet: per-channel totals and per-pair valid counts ---
    if summary is not None and not summary.empty and "Frame" in summary.columns:
        summary = _strip_managed_summary_columns(summary)
        summary_frames: set[int] = set()
        try:
            summary_frames = {
                int(v)
                for v in summary["Frame"].tolist()
                if v is not None and pd.notna(v)
            }
        except Exception:
            summary_frames = set()

        rows_by_label = {lbl: rows for lbl, rows in channel_rows}

        for label in rows_by_label:
            col_total = f"Total Foci Count ({label})"
            col_in = f"Total Foci Count within Cells ({label})"
            totals: dict[int, int] = defaultdict(int)
            within: dict[int, int] = defaultdict(int)
            for F in summary_frames:
                csv_fr = int(F) - int(frame_shift)
                totals[F] = int(raw_counts.get((label, csv_fr), 0))
            for r in rows_by_label[label] or []:
                try:
                    if not bool(r.get("in_cell", False)):
                        continue
                    cid = r.get("cell_id", None)
                    if cid in (None, "", "None") or int(cid) <= 0:
                        continue
                    fr_sum = int(r.get("frame", 0) or 0) + int(frame_shift)
                    if fr_sum in summary_frames:
                        within[fr_sum] += 1
                except Exception:
                    continue
            summary[col_total] = summary["Frame"].map(lambda f: int(totals.get(int(f), 0)) if f is not None and pd.notna(f) else 0)
            summary[col_in] = summary["Frame"].map(lambda f: int(within.get(int(f), 0)) if f is not None and pd.notna(f) else 0)

        for la, lb, df, _count in pair_tables or []:
            rows_a = rows_by_label.get(la) or []
            rows_b = rows_by_label.get(lb) or []
            col_name = _pair_column_name(la, lb)
            per_fr = _valid_pairs_counts_by_summary_frame(
                df=df,
                frame_shift=frame_shift,
                same_cell_only=same_cell_only,
                rows_a=rows_a,
                rows_b=rows_b,
                summary_frames=summary_frames,
            )
            summary[col_name] = summary["Frame"].map(
                lambda f: int(per_fr.get(int(f), 0)) if f is not None and pd.notna(f) else 0
            )

    with pd.ExcelWriter(analysis_path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        if summary is not None and not summary.empty:
            summary.to_excel(writer, sheet_name="Summary", index=False)
        per_cell.to_excel(writer, sheet_name="Per-Cell Data", index=False)


def _read_foci_csv_rows(csv_path: Path) -> list[dict[str, Any]]:
    import csv

    if csv_path is None or not csv_path.exists():
        return []
    rows: list[dict[str, Any]] = []
    with open(csv_path, "r", newline="") as f:
        reader = csv.DictReader(f)
        for r in reader:
            try:
                y = int(float(r.get("y", 0) or 0))
                x = int(float(r.get("x", 0) or 0))
            except Exception:
                continue
            try:
                intensity = float(r.get("intensity", 0) or 0.0)
            except Exception:
                intensity = 0.0
            prob_raw = r.get("probability", None)
            prob = None
            if prob_raw not in (None, "", "None"):
                try:
                    prob = float(prob_raw)
                except Exception:
                    prob = None
            frame = 0
            if "frame" in (reader.fieldnames or []):
                try:
                    frame = int(float(r.get("frame", 0) or 0))
                except Exception:
                    frame = 0
            rows.append({"frame": frame, "y": y, "x": x, "intensity": intensity, "probability": prob})
    return rows


def _csv_path_for_source(
    pipeline_dir: Path,
    *,
    mask_path: Path,
    foci_source_path: Path,
    prefer_filtered: bool,
) -> Optional[Path]:
    filt = filtered_foci_csv_path(pipeline_dir, mask_path=mask_path, foci_source_path=foci_source_path)
    if prefer_filtered and filt.exists():
        return filt
    canon = foci_csv_path(pipeline_dir, mask_path=mask_path, foci_source_path=foci_source_path)
    if canon.exists():
        return canon
    legacy = raw_foci_csv_path(pipeline_dir, mask_path=mask_path, foci_source_path=foci_source_path)
    if legacy.exists():
        return legacy
    for cand in legacy_foci_csv_candidates(pipeline_dir, mask_path=mask_path, foci_source_path=foci_source_path):
        if cand.exists():
            return cand
    return None


def _channel_label_for_source(tab, mask_path: Path, foci_source_path: Path, used: set[str]) -> str:
    mask_stem = mask_path.stem.replace("_masks", "")
    tag = None
    try:
        tag = tab._extract_channel_tag_for_mask_stem(Path(foci_source_path).stem, mask_stem)
    except Exception:
        tag = None
    base = (tag or "").strip() or "foci"
    lbl = base
    n = 2
    while lbl in used:
        lbl = f"{base}_{n}"
        n += 1
    used.add(lbl)
    return lbl


def _assigned_foci_rows_for_channel(
    mask_labels: np.ndarray,
    channel_label: str,
    csv_rows: list[dict[str, Any]],
    *,
    frame: int,
) -> list[dict[str, Any]]:
    from bactimate.modules.foci_utils import assign_peaks_to_cells

    peaks = [r for r in csv_rows if int(r.get("frame", 0) or 0) == int(frame)]
    assigned, _ = assign_peaks_to_cells(peaks, mask_labels, dilation_px=1, frame_index=int(frame))
    out: list[dict[str, Any]] = []
    for r in assigned:
        q = dict(r)
        q["channel_name"] = channel_label
        out.append(q)
    return out


def should_run_post_analysis_colocalization(tab) -> bool:
    cb = getattr(tab, "coloc_calc_all_pairs", None)
    return bool(cb is not None and cb.isChecked())


def run_post_analysis_colocalization_for_masks(tab, mask_paths: list[Path] | None) -> list[str]:
    all_lines: list[str] = []
    if not mask_paths:
        return all_lines
    if not should_run_post_analysis_colocalization(tab):
        return all_lines
    for mp in mask_paths:
        _n, lines = _run_all_pairs_for_mask(tab, mp)
        if lines:
            all_lines.extend(lines)
            for line in lines:
                _log_tab(tab, line)
    return all_lines


def _run_all_pairs_for_mask(tab, mask_path: Path) -> tuple[int, list[str]]:
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.utils.dataframe import dataframe_to_rows
    from bactimate.modules.analysis_coloc import ColocSettings, analyze_colocalization_foci_table

    log_lines: list[str] = []
    mask_path = Path(mask_path)
    log_lines.append(f"[Coloc] Start: {mask_path.name}\n")
    mask_frames = _load_mask_frames(mask_path)
    if not mask_frames:
        return 0, [f"[Coloc] Skipped {mask_path.name}: could not read mask.\n"]

    prefer_filtered = bool(getattr(tab, "use_filtered_masks_cb", None) and tab.use_filtered_masks_cb.isChecked())
    pipeline_dir = mask_path.parent
    try:
        sources = tab._resolve_foci_paths_for_mask(mask_path)
    except Exception:
        sources = []
    # Colocalization is independent from Analyze FI:
    # if no explicit *_foci sources are found, also try FI-matched stacks.
    if not sources:
        try:
            sources = tab._resolve_fi_paths_for_mask(mask_path)
        except Exception:
            sources = []
    if not sources:
        return 0, [f"[Coloc] Skipped {mask_path.name}: no matching *_foci sources.\n"]
    log_lines.append(
        f"[Coloc] Sources for {mask_path.name}: {len(sources)} -> "
        + ", ".join(Path(s).name for s in sources[:8])
        + (" ..." if len(sources) > 8 else "")
        + "\n"
    )

    used_labels: set[str] = set()
    channel_rows: list[tuple[str, list[dict[str, Any]]]] = []
    raw_by_label_frame: dict[tuple[str, int], int] = defaultdict(int)
    for src in sources:
        lbl = _channel_label_for_source(tab, mask_path, src, used_labels)
        csv_path = _csv_path_for_source(pipeline_dir, mask_path=mask_path, foci_source_path=src, prefer_filtered=prefer_filtered)
        if csv_path is None:
            log_lines.append(f"[Coloc] Source {Path(src).name} -> no matching foci CSV.\n")
            continue
        rows = _read_foci_csv_rows(csv_path)
        if not rows:
            log_lines.append(f"[Coloc] Source {Path(src).name} -> CSV {csv_path.name} has no rows.\n")
            continue
        for r in rows:
            try:
                fr_csv = int(r.get("frame", 0) or 0)
            except Exception:
                fr_csv = 0
            raw_by_label_frame[(lbl, fr_csv)] += 1
        frames = sorted({int(r.get("frame", 0) or 0) for r in rows})
        assigned_all: list[dict[str, Any]] = []
        max_mask_idx = max(mask_frames.keys()) if mask_frames else 0
        for fr in frames:
            fr_use = max(0, min(int(fr), int(max_mask_idx)))
            mask_labels = mask_frames.get(fr_use)
            if mask_labels is None:
                continue
            assigned = _assigned_foci_rows_for_channel(mask_labels, lbl, rows, frame=fr)
            assigned_all.extend(assigned)
        channel_rows.append((lbl, assigned_all))
        log_lines.append(
            f"[Coloc] Channel '{lbl}' from {Path(src).name}: csv_rows={len(rows)}, assigned_rows={len(assigned_all)}, csv={csv_path.name}\n"
        )

    stem = stem_from_mask_path(mask_path)
    analysis_path = analysis_dir_from_base(base_dir_for_mask(mask_path)) / f"{stem}_analysis.xlsx"
    if not analysis_path.exists():
        return 0, [f"[Coloc] Skipped {mask_path.name}: analysis workbook not found ({analysis_path.name}).\n"]
    log_lines.append(
        f"[Coloc] Analysis workbook target: {analysis_path}\n"
    )
    log_lines.append(
        f"[Coloc] Channels with usable assigned rows: {len(channel_rows)} -> "
        + ", ".join(f"{lbl}({len(rows)})" for lbl, rows in channel_rows)
        + "\n"
    )

    # Build per pair tables
    same_cell_only = bool(getattr(tab, "coloc_merge_same_cell_only").isChecked())
    settings = ColocSettings(
        coloc_max_distance=float(getattr(tab, "coloc_merge_max_distance").value()),
        coloc_threshold=float(getattr(tab, "coloc_merge_threshold").value()),
        coloc_filter_boundaries=same_cell_only,
    )
    pair_tables: list[tuple[str, str, pd.DataFrame, int]] = []
    if len(channel_rows) >= 2:
        for (la, rows_a), (lb, rows_b) in itertools.combinations(channel_rows, 2):
            foci_rows = []
            foci_rows.extend(rows_a)
            foci_rows.extend(rows_b)
            out = analyze_colocalization_foci_table(foci_rows, settings, foci_channels=(la, lb), frame=None)
            df = (
                pd.DataFrame(out.pair_records)
                if out.pair_records
                else pd.DataFrame({"message": ["No colocalized focus pairs."]})
            )
            count = 0
            if not ("message" in df.columns and len(df.columns) == 1):
                count = int(
                    sum(
                        1
                        for _, row in df.iterrows()
                        if bool(row.get("is_valid", True))
                        and ((not same_cell_only) or bool(row.get("same_cell", False)))
                    )
                )
            pair_tables.append((la, lb, df, count))
    else:
        log_lines.append(
            f"[Coloc] No sheet write: need >=2 channels with data, got {len(channel_rows)}.\n"
        )

    # Update Per-Cell Data, Summary foci/pair counts, and (below) Colocalization sheets.
    if channel_rows:
        try:
            _persist_workbook_count_updates(
                analysis_path=analysis_path,
                channel_rows=channel_rows,
                raw_counts=dict(raw_by_label_frame),
                pair_tables=pair_tables,
                same_cell_only=same_cell_only,
            )
        except Exception as exc:
            log_lines.append(
                f"[Coloc] Warning: could not update Per-Cell Data / Summary counts in "
                f"{analysis_path.name}: {type(exc).__name__}: {exc}\n"
            )

    if not pair_tables:
        # No coloc sheets to write, but Per-Cell counts may still have been updated.
        log_lines.append("[Coloc] No channel pairs available -> no Colocalization sheet written.\n")
        return 0, log_lines

    # Atomically replace Colocalization sheets in-memory, then save once.
    try:
        wb = load_workbook(analysis_path)
        used = set(wb.sheetnames)
        existing_coloc = [n for n in wb.sheetnames if isinstance(n, str) and n.startswith("Colocalization ")]
        log_lines.append(
            f"[Coloc] Workbook before write: total_sheets={len(wb.sheetnames)}, existing_coloc={len(existing_coloc)}.\n"
        )
        for name in list(wb.sheetnames):
            if isinstance(name, str) and name.startswith("Colocalization "):
                del wb[name]
                used.discard(name)
        written_sheet_names: list[str] = []
        for la, lb, df, _count in pair_tables:
            sheet_name = _safe_sheet_name(f"Colocalization {la}×{lb}", used)
            ws = wb.create_sheet(title=sheet_name)
            df_excel = df.copy()
            try:
                for c in list(df_excel.columns):
                    df_excel[c] = df_excel[c].map(_excel_cell_value)
            except Exception:
                pass
            for row in dataframe_to_rows(df_excel, index=False, header=True):
                ws.append(row)
            written_sheet_names.append(sheet_name)
        wb.save(analysis_path)
        wb.close()
        log_lines.append(
            f"[Coloc] Workbook saved: wrote {len(written_sheet_names)} Colocalization sheet(s): "
            + ", ".join(written_sheet_names)
            + "\n"
        )
        for la, lb, _df, count in pair_tables:
            log_lines.append(f"[Coloc] {mask_path.name} {la} × {lb} -> {count} pair(s) in {analysis_path.name}\n")
    except Exception as exc:
        return 0, [f"[Coloc] Failed {mask_path.name}: could not write sheets ({exc}).\n"]

    return len(pair_tables), log_lines
